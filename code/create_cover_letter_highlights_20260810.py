from __future__ import annotations

import json
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "outputs" / "final_submission_package_20260810" / "submission_files"
AUDIT = OUT / "audit"

TITLE = (
    "Association between sleep duration and heart disease mortality among US adults: "
    "an NHANES 2005–2018 linked mortality analysis"
)

HIGHLIGHTS = [
    "NHANES adults reporting at least 9 h/night had higher heart disease mortality.",
    "Associations persisted across lagged, imputed, and competing-risk analyses.",
    "Long sleep corresponded to a 1.01-percentage-point higher 10-year absolute risk.",
    "Health-status adjustment attenuated but did not eliminate the association.",
    "Long sleep may mark underlying health vulnerability rather than causation.",
]

BLUE = RGBColor(46, 116, 181)
INK = RGBColor(28, 39, 49)
MUTED = RGBColor(95, 103, 112)


def set_cell_borderless(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        item = borders.find(tag)
        if item is None:
            item = OxmlElement(f"w:{edge}")
            borders.append(item)
        item.set(qn("w:val"), "nil")


def set_run(run, *, size=11, bold=False, italic=False, color=INK) -> None:
    run.font.name = "Calibri"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Calibri")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Calibri")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = color


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.font.color.rgb = INK
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    if "Submission Title" not in doc.styles:
        style = doc.styles.add_style("Submission Title", WD_STYLE_TYPE.PARAGRAPH)
    else:
        style = doc.styles["Submission Title"]
    style.font.name = "Calibri"
    style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    style.font.size = Pt(16)
    style.font.bold = True
    style.font.color.rgb = BLUE
    style.paragraph_format.space_before = Pt(0)
    style.paragraph_format.space_after = Pt(4)
    style.paragraph_format.keep_with_next = True

    if "Submission Subtitle" not in doc.styles:
        style = doc.styles.add_style("Submission Subtitle", WD_STYLE_TYPE.PARAGRAPH)
    else:
        style = doc.styles["Submission Subtitle"]
    style.font.name = "Calibri"
    style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    style.font.size = Pt(10)
    style.font.color.rgb = MUTED
    style.paragraph_format.space_before = Pt(0)
    style.paragraph_format.space_after = Pt(14)
    style.paragraph_format.keep_with_next = True

    props = doc.core_properties
    props.author = ""
    props.last_modified_by = ""
    props.language = "en-US"


def add_header_block(doc: Document, title: str, subtitle: str) -> None:
    p = doc.add_paragraph(style="Submission Title")
    p.add_run(title)
    p = doc.add_paragraph(style="Submission Subtitle")
    p.add_run(subtitle)


def add_body_paragraph(doc: Document, text: str, *, after=6, keep=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.10
    p.paragraph_format.keep_with_next = keep
    run = p.add_run(text)
    set_run(run)
    return p


def build_cover_letter() -> Path:
    doc = Document()
    configure_document(doc)
    doc.core_properties.title = "Cover letter — Preventive Medicine Reports"
    doc.core_properties.subject = TITLE

    add_header_block(
        doc,
        "COVER LETTER",
        "Preventive Medicine Reports  |  Original Research Paper",
    )

    meta = doc.add_table(rows=2, cols=2)
    meta.autofit = False
    meta.columns[0].width = Inches(4.7)
    meta.columns[1].width = Inches(1.8)
    set_cell_borderless(meta)
    left = meta.cell(0, 0).paragraphs[0]
    left.paragraph_format.space_after = Pt(1)
    set_run(left.add_run("[Corresponding Author Name]"), bold=True)
    right = meta.cell(0, 1).paragraphs[0]
    right.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    right.paragraph_format.space_after = Pt(1)
    set_run(right.add_run("August 10, 2026"), color=MUTED)
    left = meta.cell(1, 0).paragraphs[0]
    left.paragraph_format.space_after = Pt(10)
    set_run(left.add_run("[Affiliation]  |  [Email address]"), size=10, color=MUTED)
    right = meta.cell(1, 1).paragraphs[0]
    right.paragraph_format.space_after = Pt(10)

    add_body_paragraph(doc, "Dear Editor,", after=8, keep=True)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(7)
    p.paragraph_format.line_spacing = 1.10
    set_run(p.add_run("Please consider the Original Research Paper, "))
    set_run(p.add_run(f'“{TITLE},”'), bold=True, italic=True)
    set_run(p.add_run(" for publication in Preventive Medicine Reports."))

    add_body_paragraph(
        doc,
        "This study analyzed 37,516 adults from the 2005–2018 National Health and Nutrition "
        "Examination Survey with linked mortality follow-up to characterize the association "
        "between habitual sleep duration and heart disease mortality. Survey-weighted "
        "cause-specific Cox models incorporated the NHANES sampling design. Nonlinear models, "
        "2- and 5-year lag analyses, multiple imputation, competing-risk analyses, and "
        "regression-standardized cumulative incidence estimates were used to examine the "
        "robustness and public health meaning of the association.",
        after=7,
    )

    add_body_paragraph(
        doc,
        "Adults reporting at least 9 h/night had higher heart disease mortality than those "
        "reporting 7–<8 h/night in the primary model (hazard ratio 1.70, 95% confidence interval "
        "1.30–2.22). The estimate was attenuated but remained elevated after adjustment for "
        "cardiometabolic conditions and prevalent cardiovascular disease (hazard ratio 1.53, "
        "95% confidence interval 1.18–1.99). The corresponding 10-year standardized cumulative "
        "incidence was 2.88% versus 1.87%, a risk difference of 1.01 percentage points. Across "
        "sensitivity analyses, the findings supported long sleep as a marker of underlying health "
        "vulnerability rather than evidence of a direct causal effect of sleep duration.",
        after=7,
    )

    add_body_paragraph(
        doc,
        "This work is well suited to Preventive Medicine Reports because it addresses a readily "
        "recognized behavioral health measure using nationally representative data and translates "
        "relative associations into absolute risk. Its deliberate separation of baseline-confounder "
        "adjustment from health-status attenuation, together with extensive reverse-causation and "
        "competing-risk analyses, offers a cautious and prevention-focused interpretation. The "
        "results may inform cardiovascular risk stratification while avoiding the unsupported "
        "conclusion that changing sleep duration would itself alter mortality risk.",
        after=7,
    )

    add_body_paragraph(
        doc,
        "Thank you for considering this manuscript. It should be of interest to readers working "
        "in preventive medicine, epidemiology, sleep health, and cardiovascular disease prevention. "
        "The opportunity to have it evaluated for peer review would be greatly appreciated.",
        after=10,
    )

    add_body_paragraph(doc, "Sincerely,", after=13, keep=True)
    p = add_body_paragraph(doc, "[Corresponding Author Name]", after=1, keep=True)
    p.runs[0].bold = True
    add_body_paragraph(doc, "[Affiliation]", after=1, keep=True)
    add_body_paragraph(doc, "[Email address]", after=0)

    path = OUT / "Cover_Letter_Preventive_Medicine_Reports_20260810.docx"
    doc.save(path)
    return path


def build_highlights() -> Path:
    doc = Document()
    configure_document(doc)
    doc.core_properties.title = "Highlights — Preventive Medicine Reports"
    doc.core_properties.subject = TITLE
    add_header_block(doc, "HIGHLIGHTS", "Preventive Medicine Reports  |  Separate submission file")

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(18)
    p.paragraph_format.line_spacing = 1.10
    set_run(p.add_run(TITLE), size=11, italic=True, color=INK)

    for item in HIGHLIGHTS:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.left_indent = Inches(0.5)
        p.paragraph_format.first_line_indent = Inches(-0.25)
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(8)
        p.paragraph_format.line_spacing = 1.167
        set_run(p.add_run(item), size=11)

    path = OUT / "Highlights_Preventive_Medicine_Reports_20260810.docx"
    doc.save(path)
    return path


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    AUDIT.mkdir(parents=True, exist_ok=True)
    cover = build_cover_letter()
    highlights = build_highlights()
    audit = {
        "journal": "Preventive Medicine Reports",
        "manuscript_title": TITLE,
        "highlight_count": len(HIGHLIGHTS),
        "maximum_allowed_characters": 85,
        "highlights": [{"text": text, "characters_including_spaces": len(text)} for text in HIGHLIGHTS],
        "all_within_limit": all(len(text) <= 85 for text in HIGHLIGHTS),
        "cover_letter_placeholders": [
            "[Corresponding Author Name]",
            "[Affiliation]",
            "[Email address]",
        ],
    }
    (AUDIT / "submission_documents_audit.json").write_text(
        json.dumps(audit, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )
    print(cover)
    print(highlights)
    print(json.dumps(audit["highlights"], ensure_ascii=False))


if __name__ == "__main__":
    main()
