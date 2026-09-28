"""Revise the manuscript after the pooled-survey and survival-model audit."""

from __future__ import annotations

import math
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches

from revise_manuscript_causal_extension_20260808 import (
    add_picture_before,
    add_supp_table,
    apply_semantic_styles,
    clean_mojibake,
    effect,
    fill_table,
    find_prefix,
    read_tsv,
    replace_prefix,
    replace_text,
    set_table_widths,
)


ROOT = Path(__file__).resolve().parents[1]
SOURCE_DIR = ROOT / "outputs" / "manuscript_causal_extension_20260808"
SOURCE_MAIN = SOURCE_DIR / "Manuscript_causal_extension_20260808.docx"
SOURCE_SUPPLEMENT = SOURCE_DIR / "Supplementary_methods_results_causal_extension_20260808.docx"
BASE = ROOT / "outputs" / "survey_validated_reanalysis_20260808"
EXT = ROOT / "outputs" / "survey_validated_extension_20260808"
OUTPUT_DIR = ROOT / "outputs" / "manuscript_survey_validated_20260808"
OUTPUT_MAIN = OUTPUT_DIR / "Manuscript_survey_validated_20260808.docx"
OUTPUT_SUPPLEMENT = OUTPUT_DIR / "Supplementary_methods_results_survey_validated_20260808.docx"


def pick(rows, **criteria):
    selected = [row for row in rows if all(row.get(key) == value for key, value in criteria.items())]
    if len(selected) != 1:
        raise ValueError((criteria, len(selected)))
    return selected[0]


def metric(rows, name):
    return float(pick(rows, metric=name)["value"])


def fmt_p(value):
    value = float(value)
    return "<0.001" if value < 0.001 else f"{value:.3f}"


def replace_picture_before_caption(document, caption_prefix: str, image_path: Path, width: float) -> None:
    caption = find_prefix(document, caption_prefix)
    previous = caption._p.getprevious()
    if previous is not None and previous.tag.endswith("}p") and previous.xpath(".//w:drawing"):
        previous.getparent().remove(previous)
    paragraph = caption.insert_paragraph_before()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    shape = paragraph.add_run().add_picture(str(image_path), width=Inches(width))
    shape._inline.docPr.set("descr", caption_prefix.rstrip("."))
    shape._inline.docPr.set("title", caption_prefix.rstrip("."))


def update_main_tables(document, base_rows, extension_rows) -> None:
    category, mi, fine_gray = base_rows
    lag, risks, contrasts = extension_rows
    category_index = {(row["model"], row["term"]): row for row in category}
    mi_index = {(row["model"], row["term"]): row for row in mi}
    fg_index = {(row["model"], row["term"]): row for row in fine_gray}
    lag_index = {(row["model"], row["term"]): row for row in lag}
    long_term = "sleep[>=9 h]"
    robustness_rows = [
        ["Free of prevalent CVD", effect(category_index[("CVD_free_Model2_confounder", long_term)]), effect(category_index[("CVD_free_Model3_health_status", long_term)]), category_index[("CVD_free_Model2_confounder", long_term)]["n"], category_index[("CVD_free_Model2_confounder", long_term)]["events"], "HR"],
        ["Exclude first 2 years", effect(lag_index[("lag2y_Model2_confounder", long_term)]), effect(lag_index[("lag2y_Model3_health_status", long_term)]), lag_index[("lag2y_Model2_confounder", long_term)]["n"], lag_index[("lag2y_Model2_confounder", long_term)]["events"], "HR"],
        ["Exclude first 5 years", effect(lag_index[("lag5y_Model2_confounder", long_term)]), effect(lag_index[("lag5y_Model3_health_status", long_term)]), lag_index[("lag5y_Model2_confounder", long_term)]["n"], lag_index[("lag5y_Model2_confounder", long_term)]["events"], "HR"],
        ["Multiple imputation", effect(mi_index[("MI_Model2_confounder", long_term)]), effect(mi_index[("MI_Model3_health_status", long_term)]), mi_index[("MI_Model2_confounder", long_term)]["n"], mi_index[("MI_Model2_confounder", long_term)]["events"], "HR"],
        ["Fine–Gray sensitivity*", effect(fg_index[("FineGray_Model2_confounder", long_term)]), effect(fg_index[("FineGray_Model3_health_status", long_term)]), fg_index[("FineGray_Model2_confounder", long_term)]["n"], fg_index[("FineGray_Model2_confounder", long_term)]["events"], "SHR"],
    ]
    fill_table(document.tables[2], ["Analysis", "Model 2", "Model 3", "N", "Heart deaths", "Measure"], robustness_rows, 7.0)
    set_table_widths(document.tables[2], [2200, 1750, 1750, 1000, 1100, 1100])

    risk_index = {(row["sleep_group"], float(row["horizon_years"])): row for row in risks}
    contrast_index = {(row["sleep_group"], float(row["horizon_years"])): row for row in contrasts}
    labels = ("<6 h", "6-<7 h", "7-<8 h", "8-<9 h", ">=9 h")
    display = {"<6 h": "<6 h", "6-<7 h": "6–<7 h", "7-<8 h": "7–<8 h", "8-<9 h": "8–<9 h", ">=9 h": "≥9 h"}
    table_rows = []
    for label in labels:
        cells = [display[label]]
        for horizon in (5.0, 10.0):
            row = risk_index[(label, horizon)]
            cells.append(f"{100*float(row['standardized_risk']):.2f}% ({100*float(row['lower_95_bootstrap']):.2f}–{100*float(row['upper_95_bootstrap']):.2f})")
        if label == "7-<8 h":
            cells.append("Reference")
        else:
            row = contrast_index[(label, 10.0)]
            cells.append(f"{float(row['risk_difference_percentage_points']):.2f} ({100*float(row['rd_lower_95_bootstrap']):.2f}–{100*float(row['rd_upper_95_bootstrap']):.2f})")
        table_rows.append(cells)
    fill_table(document.tables[3], ["Sleep duration", "5-year risk", "10-year risk", "10-year RD vs 7–<8 h, percentage points"], table_rows, 7.5)
    set_table_widths(document.tables[3], [1700, 2100, 2100, 3500])


def build_main() -> None:
    category = read_tsv(BASE / "causal_models_sleep_categories.tsv")
    mi = read_tsv(BASE / "multiple_imputation_rubin_pooled.tsv")
    fine_gray = read_tsv(BASE / "competing_risk_fine_gray.tsv")
    lag = read_tsv(EXT / "reverse_causation_lag_2y_5y.tsv")
    same_lag = read_tsv(EXT / "same_cycle_composition_lag_0y_2y_5y.tsv")
    ph = read_tsv(EXT / "proportional_hazards_time_interaction.tsv")
    risks = read_tsv(EXT / "standardized_absolute_risk_horizons.tsv")
    contrasts = read_tsv(EXT / "standardized_absolute_risk_contrasts.tsv")
    follow = read_tsv(EXT / "follow_up_and_event_summary.tsv")
    one_domain = read_tsv(EXT / "health_status_domain_one_at_a_time.tsv")
    sequential = read_tsv(EXT / "health_status_sequential_attenuation.tsv")
    alt_rcs = read_tsv(EXT / "alternative_4knot_sleep_rcs_summary.tsv")
    design = read_tsv(EXT / "survey_design_audit.tsv")
    evalue = read_tsv(EXT / "e_values_long_sleep.tsv")

    long_term = "sleep[>=9 h]"
    main_m2 = pick(category, model="Model2_confounder", term=long_term)
    main_m3 = pick(category, model="Model3_health_status", term=long_term)
    mi_m2 = pick(mi, model="MI_Model2_confounder", term=long_term)
    mi_m3 = pick(mi, model="MI_Model3_health_status", term=long_term)
    same_m2 = {int(row["lag_months"]): row for row in same_lag if row["model"].endswith("Model2_confounder") and row["term"] == long_term}
    same_m3 = {int(row["lag_months"]): row for row in same_lag if row["model"].endswith("Model3_health_status") and row["term"] == long_term}
    ph_m2 = pick(ph, model="Model2_confounder", term=long_term)
    ph_m3 = pick(ph, model="Model3_health_status", term=long_term)
    risk_long10 = pick(risks, sleep_group=">=9 h", horizon_years="10.0")
    risk_ref10 = pick(risks, sleep_group="7-<8 h", horizon_years="10.0")
    rd_long10 = pick(contrasts, sleep_group=">=9 h", horizon_years="10.0")
    siri_only = pick(one_domain, model="Model2_plus_SIRI")
    before_siri = pick(sequential, stage="E_plus_CVD")
    after_siri = pick(sequential, stage="F_plus_SIRI")
    alt_m2 = pick(alt_rcs, model="Model2_confounder")
    alt_m3 = pick(alt_rcs, model="Model3_health_status")
    ev_m2 = pick(evalue, model="Model2_confounder")

    document = Document(SOURCE_MAIN)
    clean_mojibake(document)
    replacements = {
        "Background: Long sleep": "Background: Long sleep is associated with cardiovascular mortality, but the signal may reflect underlying illness, frailty, or other unmeasured factors. We evaluated a causally motivated model hierarchy and audited the complex-survey and survival-model implementation before interpreting relative and absolute risks.",
        "Methods: We analyzed adults": "Methods: We analyzed adults aged ≥20 years in NHANES 2005–2018 linked to mortality follow-up. Cause-specific Cox models used WTMEC2YR/7 and stratified-PSU covariance with explicit cycle-nested design identifiers. We evaluated proportional hazards, same-cycle-composition 0-, 2-, and 5-year lags, multiple imputation informed by survival and survey variables, alternative spline knots, competing risks, quantitative bias contours, and regression-standardized cumulative incidence with 500 stratified PSU bootstrap replicates.",
        "Results: The cohort included": f"Results: The cohort included 37,516 adults and 1,003 heart-disease deaths; 32,880 participants, 875 heart-disease deaths, and 2,629 other-cause deaths formed the common complete-case analysis. For ≥9 versus 7–<8 h/night, the HR was {float(main_m2['HR']):.2f} (95% CI {float(main_m2['lower_95']):.2f}–{float(main_m2['upper_95']):.2f}) in Model 2 and {float(main_m3['HR']):.2f} ({float(main_m3['lower_95']):.2f}–{float(main_m3['upper_95']):.2f}) in Model 3. Explicit cycle nesting and released masked design variables gave identical estimates; weighted and unweighted estimates differed. No clear time interaction was detected (≥9-h P={float(ph_m2['p_time_interaction']):.3f} and {float(ph_m3['p_time_interaction']):.3f}). In the same early-cycle source, Model 2 HRs at 0, 2, and 5 years were {float(same_m2[0]['HR']):.2f}, {float(same_m2[24]['HR']):.2f}, and {float(same_m2[60]['HR']):.2f}. The 10-year standardized risks were {100*float(risk_long10['standardized_risk']):.2f}% and {100*float(risk_ref10['standardized_risk']):.2f}% (difference {float(rd_long10['risk_difference_percentage_points']):.2f} percentage points; 95% bootstrap interval {100*float(rd_long10['rd_lower_95_bootstrap']):.2f}–{100*float(rd_long10['rd_upper_95_bootstrap']):.2f}).",
        "Conclusions: Long sleep": "Conclusions: Long sleep was a robust mortality risk marker across survey-validated models and sensitivity analyses. Health-status attenuation, same-cycle lag results, and quantitative bias contours support interpreting the association as an indicator of underlying risk rather than a demonstrated effect of sleeping longer. Bayesian shrinkage quantified uncertainty but did not repair unmeasured confounding.",
        "Potential variable roles were assigned": "Potential variable roles were assigned for the revised, causally motivated analysis from the structures in Figure 1. Age, sex, race/ethnicity, education, continuous poverty-income ratio (PIR), marital status, and smoking were treated as baseline confounders. BMI, hypertension, diabetes, and prevalent CVD were treated as health-status variables that may reflect pre-existing disease, partial pathways, or both. Prevalent CVD was positive if any of heart failure, coronary heart disease, angina, or myocardial infarction was reported; it was negative only when all four were explicitly absent. SIRI was the primary exploratory inflammatory marker and SII the sensitivity marker. Sleep quality, sleep apnea, depression, physical activity, and frailty remained incompletely measured or unavailable across cycles.",
        "Target-trial-inspired specification": "Estimand framework and alignment of eligibility and time zero",
        "A target-trial-inspired specification": "Eligibility, exposure classification, and time zero were aligned at the NHANES examination. The contrast of interest was habitual sleep ≥9 versus 7–<8 h/night, with follow-up to heart-disease death, competing death, or administrative censoring. The primary estimand was the survey-adjusted cause-specific hazard association under Model 2; a secondary estimand was the regression-standardized 5- and 10-year cumulative incidence in the presence of competing deaths. These are associational estimands: sleep is not randomized, an intervention that changes sleep duration is not uniquely defined, and exchangeability cannot be verified.",
        "Baseline characteristics were summarized": "Baseline characteristics were summarized with the 14-year MEC weight (WTMEC2YR/7) without null-hypothesis tests [15]. All cause-specific Cox models used this weight and stratified-PSU sandwich covariance. We retained the released SDMVSTRA and SDMVPSU for audit and explicitly constructed cycle × stratum and cycle × stratum × PSU identifiers for analysis. The released masked strata did not repeat across the seven cycles, so released and explicitly nested identifiers produced identical coefficients and covariance matrices. Models 1–3 were compared on the same complete-case sample. Sleep splines used knots at 5, 6, 7, 8, and 9 h/night, with a four-knot survey-weighted-percentile specification (5, 7, 8, and 9 h) as sensitivity analysis. Proportional hazards were tested with sleep-category × log(follow-up/5 years) interactions and described in 0–5-year and >5-year periods.",
        "Reverse causation was examined": "Reverse causation was examined by excluding prevalent CVD and by 2- and 5-year landmarks. To separate lagging from calendar composition, 0-, 2-, and 5-year models were also compared within the same 2005–2014 cycle source. Missing covariates were addressed with 20 chained-equation imputations: PIR and BMI used predictive mean matching, education used multinomial logistic draws, and binary variables used logistic draws. Imputation predictors included the heart-disease event, other-cause death, follow-up, Nelson–Aalen cumulative heart-disease hazard, MEC weight, masked stratum and PSU, cycle, exposure, and covariates. Fine–Gray models were retained as a weighted stratified-PSU sandwich approximation and treated as sensitivity analyses rather than the primary complex-survey estimator. Standardized risks combined two cause-specific hazards; uncertainty used 500 PSU-with-replacement bootstrap samples within strata. E-values and continuous bias contours assessed unmeasured confounding. Elastic net and Bayesian shrinkage remained secondary and were not tuned to P values.",
        "The analytic cohort comprised": f"The analytic cohort comprised 37,516 adults and 1,003 heart-disease deaths. The common complete-case sample comprised 32,880 participants, 875 heart-disease deaths, and 2,629 other-cause deaths over {metric(follow, 'total_person_years'):,.0f} person-years. Reverse-Kaplan–Meier median follow-up was {metric(follow, 'reverse_KM_follow_up_median'):.2f} years (IQR {metric(follow, 'reverse_KM_follow_up_Q1'):.2f}–{metric(follow, 'reverse_KM_follow_up_Q3'):.2f}); the survey-weighted crude heart-disease mortality rate was {metric(follow, 'survey_weighted_heart_disease_rate'):.2f} per 1,000 person-years. Strict exclusion of prevalent CVD yielded 30,096 complete cases and 531 heart-disease deaths.",
        "On the common sample, sleep categories": f"The design audit confirmed that all refitted primary Cox models used WTMEC2YR/7 and stratified-PSU covariance. Explicit cycle nesting changed neither coefficients nor covariance because released masked strata were already unique across cycles; by contrast, the unweighted ≥9-h Model 2 HR was 1.77 rather than {float(main_m2['HR']):.2f}. On the common sample, sleep categories were associated with heart-disease mortality in Model 1 (global P<0.001), Model 2 (P=0.005), and Model 3 (P=0.033). Relative to 7–<8 h/night, the ≥9-h HR was {effect(main_m2)} in Model 2 and {effect(main_m3)} in Model 3.",
        "In the common SIRI-complete sample, sequential addition": f"On the common SIRI-complete sample, adding one domain at a time to Model 2 attenuated the long-sleep log-HR by 9.3% for BMI, 7.0% for hypertension/diabetes, 11.5% for prevalent CVD, and 7.4% for log-SIRI; these percentages are non-additive. In the cumulative sequence, the HR changed from 1.71 to 1.63, 1.60, 1.54, and 1.50. Thus the incremental change when SIRI was added after CVD was {float(before_siri['HR']):.2f}→{float(after_siri['HR']):.2f}, not the full cumulative 25.4% attenuation (Figure 3).",
        "Among participants free of prevalent CVD": f"Among participants free of prevalent CVD, the ≥9-h HR was 1.81 in Model 2 and 1.67 in Model 3. In the full-cycle landmark analyses, 2- and 5-year Model 2 HRs were 1.75 and 2.14; however, the comparison also changed cycle composition. Restricting the source to 2005–2014 cycles gave 0-, 2-, and 5-year Model 2 HRs of {float(same_m2[0]['HR']):.2f}, {float(same_m2[24]['HR']):.2f}, and {float(same_m2[60]['HR']):.2f}, and Model 3 HRs of {float(same_m3[0]['HR']):.2f}, {float(same_m3[24]['HR']):.2f}, and {float(same_m3[60]['HR']):.2f}. The ≥9-h time interaction was not significant in Model 2 (P={float(ph_m2['p_time_interaction']):.3f}) or Model 3 (P={float(ph_m3['p_time_interaction']):.3f}); global four-category interaction P values were {float(ph_m2['global_sleep_time_interaction_p']):.3f} and {float(ph_m3['global_sleep_time_interaction_p']):.3f}. Multiple-imputation HRs were {float(mi_m2['HR']):.2f} and {float(mi_m3['HR']):.2f}. The Fine–Gray weighted approximation remained supportive but secondary (Figure 4; Table 3).",
        "In Model 2, the standardized cumulative incidence": f"In Model 2, the 10-year standardized cumulative incidence was {100*float(risk_long10['standardized_risk']):.2f}% (95% PSU-bootstrap interval {100*float(risk_long10['lower_95_bootstrap']):.2f}%–{100*float(risk_long10['upper_95_bootstrap']):.2f}%) for ≥9 h and {100*float(risk_ref10['standardized_risk']):.2f}% ({100*float(risk_ref10['lower_95_bootstrap']):.2f}%–{100*float(risk_ref10['upper_95_bootstrap']):.2f}%) for 7–<8 h. The risk difference was {float(rd_long10['risk_difference_percentage_points']):.2f} percentage points ({100*float(rd_long10['rd_lower_95_bootstrap']):.2f}–{100*float(rd_long10['rd_upper_95_bootstrap']):.2f}) using 500 bootstrap replicates.",
        "Risks and risk differences are percentages": "Risks and risk differences are percentages with 95% percentile intervals from 500 stratified PSU bootstrap replicates. Regression standardization combined Model 2 cause-specific Cox models for heart-disease and other-cause death and retained competing deaths in cumulative incidence.",
        "Sleep duration was associated with SIRI": f"Sleep duration was associated with SIRI in categorical and spline models, but the magnitude was small; SII showed no overall spline association. On the same SIRI-complete sample, adding log-SIRI alone to Model 2 attenuated the long-sleep log-HR by {float(siri_only['attenuation_vs_Model2_percent']):.1f}%. In the cumulative sequence, the HR changed only from {float(before_siri['HR']):.2f} before SIRI to {float(after_siri['HR']):.2f} after SIRI. The full 25.4% reduction arose across BMI, hypertension/diabetes, CVD, and SIRI together and must not be attributed to inflammation.",
        "This reanalysis supports": f"This reanalysis supports a nonlinear association between sleep duration and heart-disease mortality, with the most reproducible elevation among adults reporting ≥9 h/night. The result was unchanged by explicit cycle nesting, showed no clear violation of proportional hazards, persisted in same-cycle landmark comparisons, and corresponded to a 10-year standardized risk difference of about one percentage point. These findings support a robust risk marker, not an intervention or validated prediction effect.",
        "The spline minimum was near": f"The spline minimum was near 7 h/night. A four-knot survey-weighted-percentile sensitivity analysis retained an overall association in Model 2 (P={fmt_p(alt_m2['overall_p'])}) and Model 3 (P={fmt_p(alt_m3['overall_p'])}); Model 2 nonlinearity remained evident (P={fmt_p(alt_m2['nonlinear_p'])}). This reduces concern that the reported shape was created solely by the clinically chosen five-knot placement.",
        "The observed shape is broadly compatible": "The observed shape is broadly compatible with prior reports of U- or J-shaped sleep–mortality associations. The 5-year landmark estimate increased, but the full-cycle comparison changed calendar composition and selected survivors. Same-cycle-composition analyses reduced this ambiguity: estimates were similar at 0 and 2 years and higher at 5 years, while formal time interactions remained imprecise and nonsignificant. These results do not establish a delayed biological effect or eliminate occult disease and frailty.",
        "Long sleep may be a risk marker": f"Long sleep may be a risk marker rather than a modifiable causal exposure. The Model 2 point and confidence-limit E-values were {float(ev_m2['E_value_point']):.2f} and {float(ev_m2['E_value_CI_bound']):.2f}. Continuous bias contours show the combinations of prevalence imbalance and confounder–mortality association required to move the point estimate to the null, avoiding an arbitrary list of discrete scenarios. Neither approach controls unmeasured confounding or addresses exposure misclassification and selection bias.",
        "The inflammatory analyses provided": f"The inflammatory analyses provided limited mechanistic support. Adding log-SIRI alone changed the Model 2 HR modestly, and adding SIRI after BMI, hypertension/diabetes, and CVD changed the HR from {float(before_siri['HR']):.2f} to {float(after_siri['HR']):.2f}. Because sleep and blood counts were obtained at the same visit, the findings are pathway triangulation and health-status attenuation, not mediation.",
        "Strengths include": "Strengths include a large nationally representative cohort, an explicit and machine-auditable survey design, common-sample model comparisons, formal PH assessment, same-cycle landmark analyses, multiple imputation informed by survival and survey variables, competing-risk standardization, 500-replicate PSU bootstrap intervals, alternative spline knots, continuous quantitative-bias contours, and skeptical posterior probabilities. The analysis hierarchy places design validity and estimand clarity before penalized prediction.",
        "Several limitations remain": "Several limitations remain substantial. Sleep was self-reported once; sleep quality, obstructive sleep apnea, depression, physical activity, occupational factors, and frailty were incompletely measured or unavailable. E-values and bias contours quantify but do not remove unmeasured confounding. Landmark analyses condition on survival, Fine–Gray survey handling is an approximation, and standardized risks depend on two proportional-hazards models despite no detected time interaction. Exposure and inflammatory markers were concurrent. Results apply to the US noninstitutionalized population and should not be interpreted as evidence that changing sleep duration changes mortality risk.",
        "In NHANES 2005–2018": "In NHANES 2005–2018, long sleep was consistently associated with heart-disease mortality after explicit validation of complex-survey weighting, pooled-cycle design identifiers, proportional hazards, missing-data handling, and calendar-matched landmark analyses. The absolute 10-year difference was about 1 percentage point. Health-status attenuation, weak incremental contribution from SIRI, and quantitative bias contours favor interpreting ≥9 h/night as a robust mortality risk marker that may reflect underlying poor health. Bayesian probabilities quantify uncertainty but do not convert the association into a causal effect.",
        "Original NHANES component": "Original NHANES component and linked mortality data are available from the National Center for Health Statistics. Analysis code and tabular outputs for the survey-design audit, main Cox and spline models, PH tests, same-cycle lags, multiple imputation, competing risks, standardized risks, quantitative bias contours, inflammation, elastic net, and Bayesian sensitivity analyses accompany this revision.",
    }
    for prefix, text in replacements.items():
        replace_prefix(document, prefix, text)

    replace_prefix(document, "Figure 3.", "Figure 3. One-at-a-time and cumulative health-status adjustment for ≥9 versus 7–<8 h/night. Panel A adds each domain separately to Model 2; Panel B shows the specified cumulative order. Percent attenuation is descriptive, order-dependent, non-additive, and not a mediated proportion.")
    replace_prefix(document, "Figure 4.", "Figure 4. Robustness of the long-sleep association. Fine–Gray rows are weighted stratified-PSU sandwich approximations and are shown as secondary subdistribution-hazard sensitivity analyses; all other rows show cause-specific hazard ratios.")
    replace_prefix(document, "Figure 5.", "Figure 5. Model 2 regression-standardized cumulative incidence of heart-disease death in the presence of competing deaths. Shading shows 95% intervals from 500 stratified PSU bootstrap replicates. Curves are associational model-based contrasts, not intervention effects.")
    replace_picture_before_caption(document, "Figure 3.", EXT / "Figure3_health_status_domain_attenuation.png", 6.8)
    replace_picture_before_caption(document, "Figure 4.", EXT / "Figure4_robustness_forest.png", 6.8)
    replace_picture_before_caption(document, "Figure 5.", EXT / "Figure5_standardized_cumulative_incidence.png", 6.8)
    update_main_tables(document, (category, mi, fine_gray), (lag, risks, contrasts))
    document.core_properties.title = "Sleep duration and heart disease mortality: survey-validated NHANES analysis"
    document.core_properties.subject = "Complex-survey Cox, PH diagnostics, same-cycle lags, and absolute risk"
    document.core_properties.keywords = "NHANES; sleep; heart disease mortality; complex survey; Cox; proportional hazards"
    apply_semantic_styles(document)
    document.save(OUTPUT_MAIN)


def build_supplement() -> None:
    design = read_tsv(EXT / "survey_design_audit.tsv")
    ph = read_tsv(EXT / "proportional_hazards_time_interaction.tsv")
    same_lag = read_tsv(EXT / "same_cycle_composition_lag_0y_2y_5y.tsv")
    one_domain = read_tsv(EXT / "health_status_domain_one_at_a_time.tsv")
    alternative = read_tsv(EXT / "alternative_4knot_sleep_rcs_summary.tsv")
    follow = read_tsv(EXT / "follow_up_and_event_summary.tsv")
    risks = read_tsv(EXT / "standardized_absolute_risk_horizons.tsv")
    contrasts = read_tsv(EXT / "standardized_absolute_risk_contrasts.tsv")
    evalue = read_tsv(EXT / "e_values_long_sleep.tsv")
    sequential = read_tsv(EXT / "health_status_sequential_attenuation.tsv")

    document = Document(SOURCE_SUPPLEMENT)
    clean_mojibake(document)
    replace_prefix(document, "S2. Variable roles", "S2. Variable roles and estimand framework")
    replace_prefix(document, "Twenty stochastic", "Twenty stochastic chained-equation datasets were created. Continuous PIR and BMI used weighted Bayesian linear prediction with predictive mean matching; education used weighted Bayesian multinomial logistic draws; binary variables used weighted logistic draws. Predictors included exposure, heart-disease event, other-cause death, follow-up, the weighted Nelson–Aalen cumulative heart-disease hazard, MEC weight, released stratum and PSU, cycle, and covariates. Each completed dataset used the same cycle-nested survey Cox model and Rubin pooling. Fine–Gray models used sampling and censoring weights with a stratified-PSU sandwich approximation; because fully design-based Fine–Gray inference is less standardized, these models were secondary sensitivity analyses.")
    replace_prefix(document, "For a k-year lag", "For a k-year lag, participants with follow-up ≤k years were excluded and k years were subtracted from the remaining follow-up. The original full-cycle 5-year landmark necessarily changed cycle composition. A matched-calendar sensitivity therefore restricted the source to 2005–2014 cycles and fitted 0-, 2-, and 5-year landmarks from the same cycle source. Landmark estimates remain conditional on surviving and remaining observable past the landmark.")
    replace_prefix(document, "All stages used the same", "All models used the same SIRI-complete sample. The cumulative sequence is order-dependent. One-at-a-time models therefore added BMI, hypertension/diabetes, prevalent CVD, or log-SIRI separately to Model 2. Percent attenuation was calculated on the log-HR scale and is descriptive, non-additive, and not a mediated proportion.")
    replace_prefix(document, "Two Model 2 survey-weighted", "Two Model 2 survey-weighted cause-specific Cox models were fitted, one for heart-disease death and one for other-cause death. Breslow baseline hazards were combined to obtain individual cumulative incidence under each sleep category and averaged over the survey-weighted covariate distribution. Uncertainty used 500 PSU-with-replacement bootstrap replicates within survey strata; both models converged in all replicates.")
    replace_prefix(document, "Intervals are percentile", "Intervals are percentile intervals from 500 stratified PSU bootstrap replicates. These are associational model-based contrasts.")
    replace_prefix(document, "For estimates above 1", "For estimates above 1, E-values used E=HR+√[HR(HR−1)] for the point estimate and confidence limit closest to 1. Because absolute risk was low, the HR was treated as an approximate risk-ratio sensitivity scale. Instead of a selected list of discrete scenarios, a continuous binary-confounder grid varied reference prevalence (10%, 20%, or 30%), the long-sleep prevalence difference (0–70 percentage points, subject to prevalence ≤95%), and the confounder–mortality risk ratio (1–6). The red contour in Supplementary Figure S1 identifies combinations that move the Model 2 point estimate to 1.")

    # Update tables that changed with the refit/bootstrap.
    sequential_rows = [[row["description"], effect(row), f"{float(row['attenuation_vs_confounder_percent']):.1f}%" if row["attenuation_vs_confounder_percent"] != "nan" else "–", row["n"], row["events"]] for row in sequential]
    fill_table(document.tables[6], ["Stage", "HR (95% CI)", "Log-HR attenuation", "N", "Deaths"], sequential_rows, 7.2)
    risk_index = {(row["sleep_group"], float(row["horizon_years"])): row for row in risks}
    risk_rows = []
    for level in ("<6 h", "6-<7 h", "7-<8 h", "8-<9 h", ">=9 h"):
        row5, row10 = risk_index[(level, 5.0)], risk_index[(level, 10.0)]
        risk_rows.append([level.replace(">=", "≥").replace("-<", "–<"), f"{100*float(row5['standardized_risk']):.2f}% ({100*float(row5['lower_95_bootstrap']):.2f}–{100*float(row5['upper_95_bootstrap']):.2f})", f"{100*float(row10['standardized_risk']):.2f}% ({100*float(row10['lower_95_bootstrap']):.2f}–{100*float(row10['upper_95_bootstrap']):.2f})", "500"])
    fill_table(document.tables[7], ["Sleep duration", "5-year risk", "10-year risk", "Bootstrap replicates"], risk_rows, 7.2)
    contrast_rows = []
    for row in contrasts:
        contrast_rows.append([row["sleep_group"].replace(">=", "≥").replace("-<", "–<"), row["horizon_years"], f"{float(row['risk_difference_percentage_points']):.2f} ({100*float(row['rd_lower_95_bootstrap']):.2f}–{100*float(row['rd_upper_95_bootstrap']):.2f})", f"{float(row['risk_ratio']):.2f} ({float(row['rr_lower_95_bootstrap']):.2f}–{float(row['rr_upper_95_bootstrap']):.2f})"])
    fill_table(document.tables[8], ["Sleep duration", "Years", "Risk difference, percentage points", "Risk ratio"], contrast_rows, 7.2)
    e_rows = [[row["model"], f"{float(row['HR']):.2f}", f"{float(row['lower_95']):.2f}", f"{float(row['E_value_point']):.2f}", f"{float(row['E_value_CI_bound']):.2f}"] for row in evalue]
    fill_table(document.tables[9], ["Model", "HR", "Lower 95% limit", "Point E-value", "CI-limit E-value"], e_rows, 7.2)

    # Replace the discrete-scenario table with the continuous bias figure.
    scenario_table = document.tables[10]
    scenario_table._element.getparent().remove(scenario_table._element)
    replace_prefix(document, "Supplementary Table S11.", "Supplementary Figure S1. Continuous quantitative-bias contours for the Model 2 long-sleep estimate")
    replace_prefix(document, "The scenario table is deliberately", "Color shows the approximate bias-adjusted HR. The red HR=1 contour is a threshold, not evidence that any named unmeasured factor lies on one side of it.")
    bias_caption = find_prefix(document, "Supplementary Figure S1.")
    add_picture_before(bias_caption, EXT / "FigureS_bias_contour.png", 6.9, "Continuous quantitative-bias contour")

    anchor = find_prefix(document, "S10. Reproducibility files")
    validation_heading = anchor.insert_paragraph_before("S10. Survey-design and survival-model validation")
    validation_heading.style = anchor.style
    replace_text(anchor, "S11. Reproducibility files")
    add_picture_before(anchor, EXT / "FigureS_PH_time_specific_HR.png", 6.4, "Time-specific hazard ratios for long sleep")
    caption = anchor.insert_paragraph_before("Supplementary Figure S2. Time-specific ≥9-h hazard ratios from sleep × log(time) models")
    caption.style = "Caption"
    note = anchor.insert_paragraph_before("No ≥9-h or global sleep-category time interaction was statistically significant. Wide intervals caution against interpreting the descriptive increase as a proven delayed effect.")

    design_rows = [[row["check"].replace("_", " "), row["value"], row["status"]] for row in design]
    table = add_supp_table(document, anchor, "Supplementary Table S11. Pooled NHANES survey-design audit", ["Check", "Value", "Status"], design_rows, "Released masked strata and stratum–PSU pairs did not repeat across cycles; explicit nesting and released identifiers were numerically identical.", size=6.9)
    set_table_widths(table, [4200, 4300, 1436])

    ph_rows = []
    for row in ph:
        if row["term"] != "sleep[>=9 h]":
            continue
        ph_rows.append(["Model 2" if row["model"] == "Model2_confounder" else "Model 3", f"{float(row['gamma']):.3f} ({float(row['gamma_se']):.3f})", fmt_p(row["p_time_interaction"]), fmt_p(row["global_sleep_time_interaction_p"]), row["n"], row["events"]])
    table = add_supp_table(document, anchor, "Supplementary Table S12. Proportional-hazards interaction tests", ["Model", "≥9-h γ (SE)", "P interaction", "Global P interaction", "N", "Deaths"], ph_rows, "γ multiplies log(follow-up months/60); the main sleep coefficient is therefore the estimated log-HR at 5 years.", size=6.9)
    set_table_widths(table, [2400, 1600, 1400, 1600, 1500, 1436])

    lag_rows = []
    for months in (0, 24, 60):
        for model in ("Model2_confounder", "Model3_health_status"):
            row = next(row for row in same_lag if int(row["lag_months"]) == months and row["model"].endswith(model) and row["term"] == "sleep[>=9 h]")
            lag_rows.append([f"{months/12:g}", "Model 2" if model.startswith("Model2") else "Model 3", effect(row), row["n"], row["events"]])
    table = add_supp_table(document, anchor, "Supplementary Table S13. Same-cycle-composition landmark analyses", ["Lag, years", "Model", "HR (95% CI)", "N", "Deaths"], lag_rows, "All analyses draw their source population from NHANES 2005–2014 cycles; post-landmark N still decreases because survival beyond the landmark is required.", size=7.0)
    set_table_widths(table, [1400, 1600, 3000, 1900, 2036])

    domain_rows = [[row["model"].replace("Model2_", "Model 2 ").replace("plus_", "+ ").replace("HTN_DM", "hypertension/diabetes"), effect(row), f"{float(row['attenuation_vs_Model2_percent']):.1f}%", row["n"], row["events"]] for row in one_domain]
    table = add_supp_table(document, anchor, "Supplementary Table S14. One-at-a-time health-domain adjustment", ["Model", "HR (95% CI)", "Log-HR attenuation", "N", "Deaths"], domain_rows, "Each domain is added separately to the same Model 2. Percentages are non-additive and not causal mediation proportions.", size=7.0)
    set_table_widths(table, [3000, 2500, 2000, 1200, 1236])

    alternative_rows = [["Model 2" if row["model"] == "Model2_confounder" else "Model 3", ", ".join(f"{float(value):g}" for value in row["knots"].split(",")), fmt_p(row["overall_p"]), fmt_p(row["nonlinear_p"]), row["n"], row["events"]] for row in alternative]
    table = add_supp_table(document, anchor, "Supplementary Table S15. Alternative four-knot sleep spline", ["Model", "Knots, h", "Overall P", "Nonlinear P", "N", "Deaths"], alternative_rows, "Knots were the survey-weighted 5th, 35th, 65th, and 95th percentiles (5, 7, 8, and 9 h).", size=7.0)
    set_table_widths(table, [2300, 1500, 1300, 1500, 1700, 1636])

    follow_rows = [[row["metric"].replace("_", " "), f"{float(row['value']):,.2f}" if any(token in row["metric"] for token in ("rate", "follow_up")) else f"{float(row['value']):,.0f}", row["unit"]] for row in follow]
    table = add_supp_table(document, anchor, "Supplementary Table S16. Follow-up, person-time, and competing deaths", ["Metric", "Value", "Unit"], follow_rows, "Reverse-Kaplan–Meier follow-up uses administrative censoring as the event and deaths as censored observations.", size=7.0)
    set_table_widths(table, [4300, 2000, 3636])

    normal_note_prefixes = (
        "Color shows the approximate",
        "Released masked strata",
        "γ multiplies",
        "All analyses draw their source population",
        "Each domain is added separately",
        "Knots were the survey-weighted",
        "Reverse-Kaplan–Meier follow-up",
    )
    for paragraph in document.paragraphs:
        if paragraph.text.startswith(normal_note_prefixes):
            paragraph.style = document.styles["Normal"]

    replace_prefix(document, "Primary causal code:", "Primary analysis code: Data processing/causal_reanalysis_20260808.py. Survey-validation extension: Data processing/survey_validated_extension_20260808.py. Figure code: Data processing/create_survey_validated_figures_20260808.py. Validated outputs are in outputs/survey_validated_reanalysis_20260808/ and outputs/survey_validated_extension_20260808/. Elastic-net results remain in outputs/elastic_net_bayesian_sensitivity_20260807/.")
    document.core_properties.title = "Supplementary methods and results: survey-validated survival analysis"
    apply_semantic_styles(document, supplement=True)
    document.save(OUTPUT_SUPPLEMENT)


def main() -> None:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    build_main()
    build_supplement()
    print(OUTPUT_MAIN)
    print(OUTPUT_SUPPLEMENT)


if __name__ == "__main__":
    main()
