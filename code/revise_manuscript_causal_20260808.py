"""Revise the manuscript around a clear causal framework and validated outputs."""

from __future__ import annotations

import csv
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from docx.text.paragraph import Paragraph


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "outputs" / "manuscript_revision_20260807" / "Manuscript_revised_elastic_net_20260807.docx"
RESULTS = ROOT / "outputs" / "causal_reanalysis_20260808"
EN_RESULTS = ROOT / "outputs" / "elastic_net_bayesian_sensitivity_20260807"
OUTPUT_DIR = ROOT / "outputs" / "manuscript_causal_revision_20260808"
OUTPUT = OUTPUT_DIR / "Manuscript_causal_revision_20260808.docx"
SUPPLEMENT = OUTPUT_DIR / "Supplementary_methods_results_20260808.docx"


def read_tsv(path: Path) -> list[dict[str, str]]:
    with path.open(encoding="utf-8-sig", newline="") as handle:
        return list(csv.DictReader(handle, delimiter="\t"))


def replace_text(paragraph, text: str) -> None:
    if paragraph.runs:
        paragraph.runs[0].text = text
        for run in list(paragraph.runs[1:]):
            run._element.getparent().remove(run._element)
    else:
        paragraph.add_run(text)


def find_prefix(document: Document, prefix: str):
    matches = [paragraph for paragraph in document.paragraphs if paragraph.text.startswith(prefix)]
    if len(matches) != 1:
        raise ValueError(f"Expected one paragraph starting {prefix!r}; found {len(matches)}")
    return matches[0]


def replace_prefix(document: Document, prefix: str, text: str) -> None:
    replace_text(find_prefix(document, prefix), text)


def insert_paragraph_after(paragraph, text: str = ""):
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    new_paragraph = Paragraph(new_p, paragraph._parent)
    if paragraph.style:
        new_paragraph.style = paragraph.style
    new_paragraph.add_run(text)
    return new_paragraph


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    tr_pr.append(repeat)


def set_cell(cell, text: str, size: float = 8.2, bold: bool = False) -> None:
    cell.text = str(text)
    for paragraph in cell.paragraphs:
        paragraph.paragraph_format.space_after = Pt(0)
        for run in paragraph.runs:
            run.font.name = "Times New Roman"
            run.font.size = Pt(size)
            run.font.bold = bold
            fonts = run._element.get_or_add_rPr().rFonts
            fonts.set(qn("w:ascii"), "Times New Roman")
            fonts.set(qn("w:hAnsi"), "Times New Roman")


def delete_row(table, row) -> None:
    table._tbl.remove(row._tr)


def effect(row: dict[str, str]) -> str:
    return f"{float(row['HR']):.2f} ({float(row['lower_95']):.2f}–{float(row['upper_95']):.2f})"


def by_model_term(rows: list[dict[str, str]]) -> dict[tuple[str, str], dict[str, str]]:
    return {(row["model"], row["term"]): row for row in rows}


def replace_table2(document: Document) -> None:
    rows = read_tsv(RESULTS / "causal_models_sleep_categories.tsv")
    summaries = {row["model"]: row for row in read_tsv(RESULTS / "causal_models_sleep_categories_summary.tsv")}
    indexed = by_model_term(rows)
    models = ("Model1_demographic", "Model2_confounder", "Model3_health_status")
    data = [["Sleep duration", "Model 1\nHR (95% CI)", "Model 2\nHR (95% CI)", "Model 3\nHR (95% CI)"]]
    for label, term in (
        ("<6 h", "sleep[<6 h]"),
        ("6–<7 h", "sleep[6-<7 h]"),
        ("7–<8 h", None),
        ("8–<9 h", "sleep[8-<9 h]"),
        ("≥9 h", "sleep[>=9 h]"),
    ):
        values = [label]
        for model in models:
            values.append("Reference" if term is None else effect(indexed[(model, term)]))
        data.append(values)
    data.append(["Global P", *[f"{float(summaries[model]['global_sleep_p']):.3f}" if float(summaries[model]["global_sleep_p"]) >= 0.001 else "<0.001" for model in models]])
    table = document.tables[1]
    while len(table.rows) < len(data):
        table.add_row()
    while len(table.rows) > len(data):
        delete_row(table, table.rows[-1])
    for i, values in enumerate(data):
        for j, value in enumerate(values):
            set_cell(table.rows[i].cells[j], value, 8.2, bold=i == 0)
    set_repeat_table_header(table.rows[0])


def add_sensitivity_table(document: Document, before_prefix: str) -> None:
    category = by_model_term(read_tsv(RESULTS / "causal_models_sleep_categories.tsv"))
    category_summary = {row["model"]: row for row in read_tsv(RESULTS / "causal_models_sleep_categories_summary.tsv")}
    mi = by_model_term(read_tsv(RESULTS / "multiple_imputation_rubin_pooled.tsv"))
    mi_summary = {row["model"]: row for row in read_tsv(RESULTS / "multiple_imputation_summary.tsv")}
    fg = by_model_term(read_tsv(RESULTS / "competing_risk_fine_gray.tsv"))
    fg_summary = {row["model"]: row for row in read_tsv(RESULTS / "competing_risk_fine_gray_summary.tsv")}
    specs = [
        ("CVD-free, Model 2", category, "CVD_free_Model2_confounder", category_summary),
        ("CVD-free, Model 3", category, "CVD_free_Model3_health_status", category_summary),
        ("20-imputation, Model 2", mi, "MI_Model2_confounder", mi_summary),
        ("20-imputation, Model 3", mi, "MI_Model3_health_status", mi_summary),
        ("Fine–Gray, Model 2", fg, "FineGray_Model2_confounder", fg_summary),
        ("Fine–Gray, Model 3", fg, "FineGray_Model3_health_status", fg_summary),
    ]
    terms = ("sleep[<6 h]", "sleep[6-<7 h]", "sleep[8-<9 h]", "sleep[>=9 h]")
    anchor = find_prefix(document, before_prefix)
    caption = anchor.insert_paragraph_before("Table 3. Reverse-causation, missing-data, and competing-risk sensitivity analyses")
    caption.style = find_prefix(document, "Table 2.").style
    table = document.add_table(rows=1, cols=6)
    table.style = document.tables[1].style
    headers = ("Analysis", "<6 h", "6–<7 h", "8–<9 h", "≥9 h", "Global P")
    for j, value in enumerate(headers):
        set_cell(table.rows[0].cells[j], value, 7.4, True)
    set_repeat_table_header(table.rows[0])
    for label, source, model, summaries in specs:
        cells = table.add_row().cells
        values = [label, *[effect(source[(model, term)]) for term in terms]]
        p_value = float(summaries[model]["global_sleep_p"])
        values.append("<0.001" if p_value < 0.001 else f"{p_value:.3f}")
        for j, value in enumerate(values):
            set_cell(cells[j], value, 7.2)
    caption._p.addnext(table._tbl)
    note = document.add_paragraph(
        "Values are hazard ratios (95% confidence intervals), except Fine–Gray rows, which are subdistribution hazard ratios. The reference is 7–<8 h/night. Model 2 adjusts the designated baseline confounders; Model 3 additionally adjusts cardiometabolic and pre-existing disease status."
    )
    note.style = find_prefix(document, "All models used").style
    table._tbl.addnext(note._p)


def add_picture_before(document: Document, caption_prefix: str, path: Path, width: float, alt_text: str) -> None:
    caption = find_prefix(document, caption_prefix)
    paragraph = caption.insert_paragraph_before()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    shape = paragraph.add_run().add_picture(str(path), width=Inches(width))
    shape._inline.docPr.set("descr", alt_text)
    shape._inline.docPr.set("title", alt_text)
    paragraph.paragraph_format.space_after = Pt(4)


def relocate_table2(document: Document) -> None:
    anchor = find_prefix(document, "Additional adjustment for cardiometabolic")
    caption = find_prefix(document, "Table 2.")
    note = find_prefix(document, "All models used")
    table = document.tables[1]
    anchor._p.addnext(caption._p)
    caption._p.addnext(table._tbl)
    table._tbl.addnext(note._p)


def clean_mojibake(document: Document) -> None:
    replacements = {"每": "–", "≡": "≥", "＊": "’", "坼": "-"}
    for paragraph in document.paragraphs:
        for run in paragraph.runs:
            for old, new in replacements.items():
                run.text = run.text.replace(old, new)
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        for old, new in replacements.items():
                            run.text = run.text.replace(old, new)


def add_references(document: Document) -> None:
    ref14 = find_prefix(document, "14.\t")
    references = [
        "15.\tNational Center for Health Statistics: NHANES survey methods and analytic guidelines. Centers for Disease Control and Prevention. https://wwwn.cdc.gov/nchs/nhanes/analyticguidelines.aspx (accessed 8 August 2026).",
        "16.\tWhite IR, Royston P, Wood AM: Multiple imputation using chained equations: issues and guidance for practice. Stat Med 2011, 30:377–399.",
        "17.\tFine JP, Gray RJ: A proportional hazards model for the subdistribution of a competing risk. J Am Stat Assoc 1999, 94:496–509.",
        "18.\tVanderWeele TJ: Causal mediation analysis with survival data. Epidemiology 2011, 22:582–585.",
        "19.\tSchuler MS, Coffman DL, Stuart EA, Nguyen TQ, Vegetabile B, McCaffrey DF: Practical challenges in mediation analysis: a guide for applied researchers. Health Serv Outcomes Res Methodol 2025, 25:57–84.",
        "20.\tGeorgeson AR, Alvarez-Bartolo D, MacKinnon DP: A sensitivity analysis for temporal bias in cross-sectional mediation. Psychol Methods 2025, 30:1326–1344.",
    ]
    current = ref14
    for text in references:
        current = insert_paragraph_after(current, text)
        current.style = ref14.style


def build_main_manuscript() -> None:
    document = Document(SOURCE)
    clean_mojibake(document)
    replacements = {
        "Background: Sleep duration": "Background: Sleep duration is associated with cardiovascular outcomes, but long sleep may reflect underlying illness rather than a causal exposure. We examined whether the sleep–heart disease mortality association persisted across causal-model definitions, reverse-causation checks, missing-data analysis, competing risks, and an exploratory inflammatory pathway.",
        "Methods: We analyzed adults": "Methods: We analyzed adults aged ≥20 years in NHANES 2005–2018 linked to mortality follow-up. Sleep was grouped as <6, 6–<7, 7–<8 (reference), 8–<9, and ≥9 h/night. Survey-weighted cause-specific Cox models used flexible restricted cubic splines for age, poverty-income ratio, and BMI. Baseline confounders were separated from cardiometabolic and pre-existing disease status. Core analyses included restricted cubic splines for sleep, a baseline CVD-free cohort, a 24-month lag, 20-fold multiple imputation with Rubin pooling, Fine–Gray competing-risk models, and exploratory SIRI/SII pathway analyses.",
        "Results: The rebuilt cohort": "Results: The cohort included 37,516 adults and 1,003 heart disease deaths; the common complete-case sample included 32,880 adults and 875 deaths. For ≥9 versus 7–<8 h/night, the HR was 1.70 (95% CI 1.30–2.22) in the main confounder model and 1.53 (1.18–1.99) after additional health-status adjustment. The main spline was nonlinear (P-overall=0.006; P-nonlinear=0.006). Among 30,096 participants free of CVD at baseline (531 deaths), the corresponding HRs were 1.81 (1.28–2.57) and 1.67 (1.19–2.36). Results persisted after 20-fold imputation and in Fine–Gray models. Sleep duration was weakly and nonlinearly associated with SIRI but not SII; these contemporaneous measurements did not support a causal mediation estimate.",
        "Conclusions: After correcting": "Conclusions: Sleep duration showed a nonlinear association with heart disease mortality, with the most reproducible excess risk among adults reporting ≥9 h/night. Attenuation after health-status adjustment, together with persistence in CVD-free and lag analyses, supports interpreting long sleep as a robust risk marker that may partly reflect underlying health. Inflammatory findings were modest and exploratory. Bayesian shrinkage and elastic net quantified stability but did not resolve confounding or reverse causation.",
        "Keywords: sleep duration": "Keywords: sleep duration; heart disease mortality; NHANES; complex survey; restricted cubic spline; multiple imputation; competing risks; systemic inflammation",
        "Long sleep is especially difficult": "Long sleep is especially difficult to interpret because it may reflect poor sleep quality, depression, low physical activity, socioeconomic disadvantage, sleep-disordered breathing, or occult disease rather than a direct harmful exposure [8, 9, 11–14]. Short sleep has stronger experimental links to sympathetic and inflammatory activation [10], creating a potential mismatch between the best-established mechanism and the strongest epidemiologic signal. We therefore asked whether the nonlinear sleep–heart disease mortality association, particularly for ≥9 h/night, persisted after separating confounder control from health-status adjustment, excluding prevalent CVD, addressing missing data and competing risks, and examining SIRI and SII as exploratory inflammatory pathway markers.",
        "This observational cohort study combined": "This observational cohort study combined seven NHANES cycles from 2005–2006 through 2017–2018 and linked participants to public-use mortality follow-up. NHANES uses a multistage probability design to represent the noninstitutionalized US population [15]. Adults with eligible mortality follow-up and sleep duration from 3 to 11 h/night formed the analytic cohort (n=37,516). SIRI and SII were not required for cohort entry because they were not primary adjustment variables. Cross-cycle harmonization and quality-control procedures are detailed in the Supplementary Methods.",
        "Figure 1. Flow diagram": "Figure 1. Causal structures guiding model interpretation. Panel A shows the etiologic pathway of interest; Panel B shows an alternative in which underlying poor health causes both longer sleep and mortality. Solid arrows indicate assumed causal relations and dashed arrows mark temporally uncertain pathways. SIRI and SII were measured at the same baseline visit as sleep duration; therefore, pathway analyses are exploratory rather than causal mediation estimates.",
        "Usual sleep duration was self-reported": "Usual sleep duration was self-reported in hours per night and harmonized across cycles. It was analyzed in five clinically interpretable categories: <6, 6–<7, 7–<8, 8–<9, and ≥9 h/night, with 7–<8 h/night as the reference, and continuously with restricted cubic splines. The primary outcome was heart disease mortality, defined as MORTSTAT=1 and UCOD_LEADING=1. Other-cause deaths were censored in cause-specific Cox models and treated as competing events in Fine–Gray models.",
        "For the corrected analysis": "Potential variable roles were assigned before refitting models from the causal structures in Figure 1. Age, sex, race/ethnicity, education, continuous poverty-income ratio (PIR), marital status, and smoking were treated as baseline confounders. BMI, hypertension, diabetes, and prevalent CVD were treated as cardiometabolic or health-status variables that could represent pre-existing disease, partial pathways, or both. Prevalent CVD was positive if any of heart failure, coronary heart disease, angina, or myocardial infarction was reported; it was negative only when all four were explicitly absent. SIRI was the primary exploratory inflammatory marker and SII the sensitivity marker.",
        "Covariate selection with LASSO": "Causal framework and model definitions",
        "Penalized selection was not used": "Model 1 adjusted for age, sex, and race/ethnicity. Model 2, the main confounder model, additionally adjusted for education, continuous PIR, marital status, and smoking. Model 3, the health-status model, additionally adjusted for continuous BMI, hypertension, diabetes, and prevalent CVD. Age, PIR, and BMI were modeled with four-knot restricted cubic splines at survey-weighted 5th, 35th, 65th, and 95th percentiles. Elastic net was not used to choose confounders; its prediction-oriented results were moved to the Supplement.",
        "Baseline characteristics were summarized": "Baseline characteristics were summarized with 14-year MEC weights (WTMEC2YR/7) without null-hypothesis tests. Cause-specific Cox models incorporated the MEC weights, SDMVPSU, and SDMVSTRA, with stratified-PSU robust covariance. Models 1–3 were compared in the same complete-case sample. The four nonreference sleep-category coefficients were tested jointly. Sleep splines used knots at 5, 6, 7, 8, and 9 h/night with 7 h/night as the reference. The primary estimand was an adjusted association rather than an identified causal effect.",
        "Restricted cubic splines used knots": "Reverse causation was examined by excluding prevalent CVD and by starting follow-up after 24 months. Missing covariates were addressed with 20 stochastic chained-equation imputations: predictive mean matching for continuous variables and education, and weighted logistic models for binary variables. Imputation models included exposure, outcome status, follow-up, cycle, and other covariates; each imputed dataset was analyzed with the survey-weighted Cox model and combined using Rubin’s rules [16]. Fine–Gray models treated non-heart-disease deaths as competing events [17]. For pathway triangulation, survey-weighted linear models related sleep categories and splines to log-SIRI and log-SII; cause-specific Cox models then examined each marker with sleep in the same sample. No proportion mediated was estimated because sleep and inflammation were measured contemporaneously and Cox coefficient attenuation is not a causal indirect effect [18–20]. Elastic-net and Bayesian shrinkage analyses were restricted to the Supplement and were not tuned to P values.",
        "The rebuilt analytic cohort": "The analytic cohort comprised 37,516 adults and 1,003 heart disease deaths. The common complete-case sample comprised 32,880 participants and 875 deaths. Strict exclusion of prevalent CVD yielded 30,096 complete cases and 531 heart disease deaths. SIRI was positive and observed for 35,583 participants and SII for 35,584. Baseline characteristics are shown in Table 1; item missingness is reported rather than concealed by deterministic single-value imputation.",
        "LASSO covariate selection": "Main Cox models",
        "In the secondary weighted Cox elastic-net": "On the common sample, sleep categories were associated with heart disease mortality in Model 1 (global P<0.001), the main confounder Model 2 (P=0.005), and the health-status Model 3 (P=0.033). Relative to 7–<8 h/night, the ≥9-h HR was 1.70 (95% CI 1.30–2.22) in Model 2 and attenuated to 1.53 (1.18–1.99) in Model 3. In Model 2, estimates for <6 h and 6–<7 h were also elevated, but both attenuated after health-status adjustment; the 8–<9-h interval included 1 in both models.",
        "On the same elastic-net complete-case sample": "Additional adjustment for cardiometabolic and pre-existing disease status therefore reduced the ≥9-h log-HR by about one fifth without reversing its direction. This pattern is consistent with partial confounding, pathway blocking, or both; it should not be interpreted as a formal decomposition of total and direct causal effects.",
        "Figure 2. Secondary weighted Cox": "Figure 2. Survey-weighted restricted cubic splines for sleep duration and heart disease mortality in the full analytic cohort (Panel A) and among participants free of prevalent CVD at baseline (Panel B). Both panels use the main confounder model, 7 h/night as the reference, and knots at 5, 6, 7, 8, and 9 h/night. Shaded regions are 95% confidence intervals.",
        "Cox regression": "Reverse-causation, missing-data, and competing-risk analyses",
        "On the common complete-case sample, the unadjusted": "Among participants free of prevalent CVD, the ≥9-h HR was 1.81 (95% CI 1.28–2.57) in Model 2 and 1.67 (1.19–2.36) in Model 3; the Model 2 sleep spline remained nonlinear (P-overall=0.027; P-nonlinear=0.033). After excluding the first 24 months, the corresponding ≥9-h HRs were 1.75 (1.33–2.29) and 1.57 (1.20–2.05). With 20-fold multiple imputation, the ≥9-h HRs were 1.73 (1.35–2.23) and 1.58 (1.24–2.02). Fine–Gray subdistribution HRs were 1.53 (1.17–2.02) and 1.39 (1.06–1.83). The direction was therefore stable, while magnitude depended on whether health-status variables were controlled.",
        "Table 2. Cox models": "Table 2. Survey-weighted Cox models for sleep duration and heart disease mortality",
        "All three models used": "All models used the same 32,880 participants and 875 heart disease deaths. Model 1 adjusted for age, sex, and race/ethnicity. Model 2 additionally adjusted for education, continuous PIR, marital status, and smoking. Model 3 additionally adjusted for continuous BMI, hypertension, diabetes, and prevalent CVD. Age, PIR, and BMI were modeled with restricted cubic splines. The reference was 7–<8 h/night; global P jointly tested the four nonreference categories.",
        "Restricted cubic spline analysis": "Exploratory inflammatory pathway analyses",
        "The survey-weighted spline showed": "Sleep duration was associated with SIRI in categorical (global P=0.008) and spline models (P-overall=0.009; P-nonlinear=0.004), but the magnitude was small: adjusted geometric means were 3.4% higher for <6 h and 3.9% higher for ≥9 h than for 7–<8 h. SII showed no overall spline association (P=0.238) and no elevation for ≥9 h. On the SIRI-complete sample, the ≥9-h mortality HR changed from 1.69 (1.29–2.21) to 1.63 (1.24–2.13) after adding log-SIRI; log-SIRI itself was associated with mortality. This attenuation was not interpreted as a mediated proportion because the marker and exposure were contemporaneous and the hazard ratio is non-collapsible.",
        "Figure 3. Survey-weighted restricted cubic spline": "Figure 3. Survey-weighted restricted cubic splines relating sleep duration to SIRI (Panel A) and SII (Panel B). Values are ratios of adjusted geometric means relative to 7 h/night; shaded regions are 95% confidence intervals. Models adjust the designated baseline confounders. Concurrent measurement of sleep and inflammation precludes a strong causal-mediation interpretation.",
        "Bayesian sensitivity analysis": "Secondary penalized and Bayesian analyses",
        "The design-based Model 3 log-HR": "The elastic-net one-standard-error rule selected no additional laboratory or health-status variable in the primary or repeated PSU-clustered folds. The prediction-optimal lambda-min rule selected eight variables and attenuated the ≥9-h HR on its smaller complete-case sample. In the present health-status model, a skeptical Normal(0, 0.20²) prior produced a posterior HR of 1.34 (95% credible interval 1.08–1.67) and P(HR>1)=0.996 for ≥9 h. These analyses support directional stability but cannot correct exposure error, unmeasured confounding, or reverse causation; full methods and results are in the Supplement.",
        "After rebuilding the cohort": "This reanalysis supports a nonlinear association between sleep duration and heart disease mortality, with the most reproducible elevation among adults reporting ≥9 h/night. The ≥9-h estimate was 1.70 in the main confounder model and fell to 1.53 after cardiometabolic and pre-existing disease status were added. Persistence after CVD exclusion, a 24-month lag, multiple imputation, and competing-risk analysis argues against a result created solely by one modeling choice, but attenuation across health-status models cautions against a direct causal interpretation.",
        "The difference between the null continuous": "The spline minimum was near 7 h/night in both the full and CVD-free cohorts. A single linear term would average opposing local slopes and is therefore an inadequate summary of this pattern. Flexible adjustment for age, PIR, and BMI also reduced residual-confounding concerns created by assuming linearity or by coarse categorization.",
        "The observed shape is broadly compatible": "The observed shape is broadly compatible with prior reports of U- or J-shaped sleep–mortality associations [5–7]. However, the estimates for short and moderately long sleep were more sensitive to the model definition than the ≥9-h estimate. The CVD-free analysis did not eliminate the long-sleep association, but its wider interval and continued health-status attenuation show that absence of diagnosed CVD is not equivalent to absence of occult disease or frailty.",
        "Long sleep may be a risk marker": "Long sleep may be a risk marker rather than a modifiable causal exposure. It can accompany fragmented sleep, depression, low activity, unemployment, chronic illness, frailty, or sleep-disordered breathing [8, 9, 11–14]. The CVD-free and 24-month-lag results reduce one form of reverse causation but do not address undiagnosed disease, time-varying health, or unmeasured sleep quality. Bayesian posterior probabilities describe uncertainty conditional on the chosen likelihood and prior; they do not repair these sources of bias.",
        "Self-reported long sleep may help": "The inflammatory analyses provided only limited mechanistic support. SIRI showed a shallow nonlinear relation with sleep, whereas SII did not, and adding SIRI produced modest attenuation of the ≥9-h coefficient. Because sleep and blood counts were obtained at the same visit, chronic illness could influence both. The results are best described as pathway triangulation: systemic inflammation is associated with subsequent heart disease mortality, but the data do not establish that it mediates the sleep association.",
        "Strengths include a large national cohort": "Strengths include a large nationally representative cohort, linked cause-specific mortality, explicit separation of confounders from health-status variables, flexible continuous adjustment, common-sample comparisons, baseline CVD exclusion, lag analysis, 20-fold multiple imputation, competing-risk regression, and transparent inflammatory pathway analyses. The analysis also shows why prediction-oriented elastic net and Bayesian shrinkage should remain secondary in an etiologic study.",
        "Several limitations remain substantial": "Several limitations remain substantial. Sleep duration was self-reported once, and sleep quality, obstructive sleep apnea, depression, physical activity, occupational factors, and frailty were incompletely measured or unavailable in the harmonized data. Exposure and inflammatory markers were concurrent, precluding temporal identification of a natural or interventional indirect effect. Multiple imputation assumes the specified conditional models are adequate and that missingness is explainable by observed data. The Fine–Gray model addresses competing deaths but not residual confounding. The 3–11-h restriction and proportional-hazards structures remain modeling assumptions. Results are associations in the US noninstitutionalized population and should not be interpreted as evidence that changing sleep duration will change mortality risk.",
        "In NHANES 2005-2018": "In NHANES 2005–2018, sleep duration had a nonlinear association with heart disease mortality, concentrated most consistently among adults reporting ≥9 h/night. The association persisted in baseline CVD-free, lagged, imputed, and competing-risk analyses but attenuated with additional health-status adjustment. Long sleep should therefore be interpreted as a robust risk marker that may partly reflect underlying poor health. SIRI offered only modest exploratory pathway evidence, and neither elastic net nor Bayesian shrinkage converted the finding into a causal conclusion.",
        "Original NHANES component": "Original NHANES component and linked mortality data are available from the National Center for Health Statistics. Analysis code and tabular outputs for the causal-model, multiple-imputation, competing-risk, inflammatory-pathway, elastic-net, and Bayesian analyses accompany this revision.",
    }
    for prefix, text in replacements.items():
        replace_prefix(document, prefix, text)

    replace_table2(document)
    relocate_table2(document)
    add_sensitivity_table(document, "Exploratory inflammatory pathway analyses")
    add_picture_before(
        document, "Figure 1.", RESULTS / "Figure1_causal_dag.png", 5.65,
        "Two-panel causal diagram contrasting the etiologic sleep-to-inflammation pathway with reverse causation from underlying poor health.",
    )
    add_picture_before(
        document, "Figure 2.", RESULTS / "Figure2_sleep_mortality_rcs.png", 6.5,
        "Restricted cubic spline hazard-ratio curves for sleep duration in the full cohort and the baseline CVD-free cohort.",
    )
    add_picture_before(
        document, "Figure 3.", RESULTS / "Figure3_inflammation_pathway_rcs.png", 6.5,
        "Restricted cubic spline geometric-mean-ratio curves for SIRI and SII across sleep duration.",
    )
    add_references(document)
    document.core_properties.title = "Sleep duration and heart disease mortality among US adults"
    document.core_properties.subject = "Causal-structure reanalysis with CVD-free, multiple-imputation, competing-risk, and inflammatory-pathway analyses"
    document.core_properties.comments = "Causal-model revision generated from validated outputs on 2026-08-08."
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)


def add_simple_table(document: Document, headers: list[str], rows: list[list[str]], font_size: float = 8.5) -> None:
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for j, value in enumerate(headers):
        set_cell(table.rows[0].cells[j], value, font_size, True)
    set_repeat_table_header(table.rows[0])
    for values in rows:
        cells = table.add_row().cells
        for j, value in enumerate(values):
            set_cell(cells[j], value, font_size)
    document.add_paragraph()


def build_supplement() -> None:
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)
    normal = document.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(10.5)
    normal_fonts = normal._element.get_or_add_rPr().rFonts
    normal_fonts.set(qn("w:ascii"), "Times New Roman")
    normal_fonts.set(qn("w:hAnsi"), "Times New Roman")
    title = document.add_heading("Supplementary methods and results", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = document.add_paragraph("Sleep duration and heart disease mortality among US adults: an NHANES 2005–2018 linked mortality analysis")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    document.add_heading("S1. Data harmonization and quality control", level=1)
    document.add_paragraph(
        "The adult cohort was reconstructed before any SIRI/SII complete-case restriction. Sleep was restricted to 3–11 h/night, mortality required positive linked follow-up, and seven-cycle MEC weights were WTMEC2YR/7. Two deterministic cross-cycle mappings were repaired from official component files: 2015–2016 smoking and 2017–2018 marital status. Prevalent CVD used strict questionnaire logic (any component yes; all four no; otherwise missing). The primary cohort included 37,516 participants and 1,003 heart disease deaths. Corrected systolic pressure was used only in the secondary elastic-net analysis because upstream SBP/DBP labels were reversed."
    )
    document.add_heading("S2. Variable roles and estimands", level=1)
    add_simple_table(
        document,
        ["Role", "Variables", "Use"],
        [
            ["Baseline confounders", "Age, sex, race/ethnicity, education, PIR, marital status, smoking", "Adjusted in Model 2; target is an adjusted association conditional on these measured common causes."],
            ["Health-status / possible pathway variables", "BMI, hypertension, diabetes, prevalent CVD", "Added in Model 3 to assess attenuation; not called a fully adjusted causal model."],
            ["Exploratory pathway markers", "log-SIRI (primary), log-SII (sensitivity)", "Modeled as contemporaneous inflammatory correlates; no causal proportion mediated estimated."],
            ["Unmeasured or incomplete factors", "Sleep quality, apnea, depression, physical activity, frailty, occult disease", "Remain potential sources of residual confounding and reverse causation."],
        ],
    )

    document.add_heading("S3. Multiple imputation and competing risks", level=1)
    document.add_paragraph(
        "Twenty stochastic chained-equation datasets were created. Continuous PIR and BMI used weighted Bayesian linear prediction with predictive mean matching; education used donor-based predictive mean matching; binary variables used weighted logistic draws. Exposure, event status, follow-up, cycle, and other covariates were included. Each completed dataset was analyzed with the same survey-weighted Cox model, and coefficients and design-based covariance matrices were combined using Rubin’s rules. Fine–Gray models retained other-cause deaths in the subdistribution risk set with censoring-survival weights and used a stratified-PSU sandwich covariance approximation."
    )
    sensitivity_rows = read_tsv(RESULTS / "multiple_imputation_rubin_pooled.tsv")
    add_simple_table(
        document,
        ["Model", "Sleep category", "HR (95% CI)", "FMI"],
        [[row["model"], row["term"].replace("sleep[", "").rstrip("]"), effect(row), f"{float(row['fraction_missing_information']):.3f}"] for row in sensitivity_rows],
        8.0,
    )

    document.add_heading("S4. Elastic-net analysis", level=1)
    document.add_paragraph(
        "Elastic net was retained only as prediction-oriented sensitivity analysis. Sleep indicators and the causal core were unpenalized; BUN, hemoglobin, HDL, total cholesterol, log-SIRI, corrected systolic blood pressure, uric acid, log-creatinine, and stroke were candidate penalized variables. Five-fold validation kept PSUs intact. Alpha values 0.25, 0.50, and 1.00 were statistically indistinguishable within one standard error, so alpha=0.50 was used for stability summaries. The lambda.1se rule selected no optional variable in the primary or five repeated fold assignments. Lambda.min selected BUN, hemoglobin, HDL, total cholesterol, log-SIRI, corrected systolic pressure, uric acid, and stroke. These results do not define a confounder set."
    )
    en_rows = read_tsv(EN_RESULTS / "elastic_net_survey_cox_refit_summary.tsv")
    add_simple_table(
        document,
        ["Refit", "N", "Deaths", "Additional variables"],
        [[row["model"], row["n"], row["events"], row["additional_variables"]] for row in en_rows],
        8.0,
    )

    document.add_heading("S5. Bayesian shrinkage sensitivity", level=1)
    document.add_paragraph(
        "The design-based log-HR and robust standard error were treated as an approximate normal likelihood and combined with zero-centered normal priors. This analysis quantifies how strongly the data update skeptical priors; it is not a full Bayesian survey-survival model and does not correct confounding, reverse causation, selection, or exposure misclassification. Priors were never chosen to maximize significance."
    )
    bayes = read_tsv(RESULTS / "bayesian_shrinkage_sensitivity.tsv")
    add_simple_table(
        document,
        ["Source model", "Prior on log(HR)", "Posterior HR (95% CrI)", "P(HR>1)"],
        [[row["source_model"], row["prior"], f"{float(row['posterior_HR']):.2f} ({float(row['credible_lower_95']):.2f}–{float(row['credible_upper_95']):.2f})", f"{float(row['probability_HR_gt_1']):.3f}"] for row in bayes],
        8.0,
    )

    document.add_heading("S6. Reproducibility files", level=1)
    document.add_paragraph(
        "Primary code: 数据处理/causal_reanalysis_20260808.py. Figure code: 数据处理/create_causal_figures_20260808.py. Validated tabular outputs are in outputs/causal_reanalysis_20260808/. Elastic-net validation outputs are in outputs/elastic_net_bayesian_sensitivity_20260807/. The original Manuscript.docx and prior analyses were not overwritten."
    )
    document.core_properties.title = "Supplementary methods and results"
    document.save(SUPPLEMENT)


def main() -> None:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    build_main_manuscript()
    build_supplement()
    print(OUTPUT)
    print(SUPPLEMENT)


if __name__ == "__main__":
    main()
