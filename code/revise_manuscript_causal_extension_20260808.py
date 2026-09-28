"""Extend the causal manuscript with absolute risks and quantitative sensitivity analyses."""

from __future__ import annotations

import csv
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, Twips
from docx.text.paragraph import Paragraph


ROOT = Path(__file__).resolve().parents[1]
SOURCE_DIR = ROOT / "outputs" / "manuscript_causal_revision_20260808"
SOURCE_MAIN = SOURCE_DIR / "Manuscript_causal_revision_20260808.docx"
SOURCE_SUPPLEMENT = SOURCE_DIR / "Supplementary_methods_results_20260808.docx"
BASE = ROOT / "outputs" / "causal_reanalysis_20260808"
EXT = ROOT / "outputs" / "causal_extension_20260808"
OUTPUT_DIR = ROOT / "outputs" / "manuscript_causal_extension_20260808"
OUTPUT_MAIN = OUTPUT_DIR / "Manuscript_causal_extension_20260808.docx"
OUTPUT_SUPPLEMENT = OUTPUT_DIR / "Supplementary_methods_results_causal_extension_20260808.docx"


def read_tsv(path: Path) -> list[dict[str, str]]:
    with path.open(encoding="utf-8-sig", newline="") as handle:
        return list(csv.DictReader(handle, delimiter="\t"))


def find_prefix(document: Document, prefix: str) -> Paragraph:
    matches = [paragraph for paragraph in document.paragraphs if paragraph.text.startswith(prefix)]
    if len(matches) != 1:
        raise ValueError(f"Expected one paragraph starting {prefix!r}; found {len(matches)}")
    return matches[0]


def replace_text(paragraph: Paragraph, text: str) -> None:
    if paragraph.runs:
        paragraph.runs[0].text = text
        for run in list(paragraph.runs[1:]):
            run._element.getparent().remove(run._element)
    else:
        paragraph.add_run(text)


def replace_prefix(document: Document, prefix: str, text: str) -> None:
    replace_text(find_prefix(document, prefix), text)


def clean_mojibake(document: Document) -> None:
    replacements = {
        "每": "–",
        "≡": "≥",
        "坼": "-",
        "＊": "’",
        "杅擂揭燴": "Data processing",
        "虏": "²",
    }
    for paragraph in document.paragraphs:
        for run in paragraph.runs:
            original = run.text
            revised = original
            for old, new in replacements.items():
                revised = revised.replace(old, new)
            if revised != original:
                run.text = revised
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        original = run.text
                        revised = original
                        for old, new in replacements.items():
                            revised = revised.replace(old, new)
                        if revised != original:
                            run.text = revised


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    tr_pr.append(repeat)


def set_cell(cell, text: str, size: float = 8.0, bold: bool = False) -> None:
    cell.text = str(text)
    for paragraph in cell.paragraphs:
        paragraph.paragraph_format.space_after = Pt(0)
        for run in paragraph.runs:
            run.font.name = "Times New Roman"
            run.font.size = Pt(size)
            run.font.bold = bold
            fonts = run._element.get_or_add_rPr().rFonts
            fonts.set(qn("w:ascii"), "Times New Roman")
            fonts.set(qn("w:hAnsi"), "Times New Roman")


def clear_table(table) -> None:
    while len(table.rows) > 1:
        table._tbl.remove(table.rows[-1]._tr)


def fill_table(table, headers, rows, size=8.0) -> None:
    clear_table(table)
    while len(table.columns) < len(headers):
        table.add_column(Inches(1))
    while len(table.columns) > len(headers):
        for row in table.rows:
            row._tr.remove(row.cells[-1]._tc)
    for j, value in enumerate(headers):
        set_cell(table.rows[0].cells[j], value, size, True)
    set_repeat_table_header(table.rows[0])
    for values in rows:
        cells = table.add_row().cells
        for j, value in enumerate(values):
            set_cell(cells[j], value, size)


def set_table_widths(table, widths_twips) -> None:
    if len(widths_twips) != len(table.columns):
        raise ValueError((len(widths_twips), len(table.columns)))
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.insert(0, tbl_w)
    tbl_w.set(qn("w:type"), "dxa")
    tbl_w.set(qn("w:w"), str(sum(widths_twips)))
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:type"), "dxa")
    # Match Word's default 120-DXA start-cell margin so visible borders align.
    tbl_ind.set(qn("w:w"), "120")
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")
    grid_columns = table._tbl.tblGrid.gridCol_lst
    for grid_col, width in zip(grid_columns, widths_twips):
        grid_col.set(qn("w:w"), str(width))
    for row in table.rows:
        for cell, width in zip(row.cells, widths_twips):
            cell.width = Twips(width)
            tc_w = cell._tc.get_or_add_tcPr().get_or_add_tcW()
            tc_w.set(qn("w:type"), "dxa")
            tc_w.set(qn("w:w"), str(width))


def insert_paragraph_after(paragraph: Paragraph, text: str = "") -> Paragraph:
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    new_paragraph = Paragraph(new_p, paragraph._parent)
    new_paragraph.style = paragraph.style
    if text:
        new_paragraph.add_run(text)
    return new_paragraph


def make_heading_before(anchor: Paragraph, text: str) -> Paragraph:
    paragraph = anchor.insert_paragraph_before()
    paragraph.style = anchor.style
    run = paragraph.add_run(text)
    run.bold = True
    run.font.name = "Times New Roman"
    run.font.size = Pt(11.5)
    return paragraph


def add_picture_before(anchor: Paragraph, path: Path, width: float, alt_text: str) -> Paragraph:
    paragraph = anchor.insert_paragraph_before()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    shape = paragraph.add_run().add_picture(str(path), width=Inches(width))
    shape._inline.docPr.set("descr", alt_text)
    shape._inline.docPr.set("title", alt_text)
    paragraph.paragraph_format.space_after = Pt(4)
    return paragraph


def add_caption_before(anchor: Paragraph, text: str, template: Paragraph) -> Paragraph:
    paragraph = anchor.insert_paragraph_before(text)
    paragraph.style = template.style
    return paragraph


def add_table_before(document: Document, anchor: Paragraph, caption_text: str, headers, rows, note_text: str, cols=None, size=7.8):
    template_caption = find_prefix(document, "Table 2.") if any(p.text.startswith("Table 2.") for p in document.paragraphs) else anchor
    caption = anchor.insert_paragraph_before(caption_text)
    caption.style = template_caption.style
    table = document.add_table(rows=1, cols=cols or len(headers))
    if document.tables:
        table.style = document.tables[0].style
    fill_table(table, headers, rows, size)
    caption._p.addnext(table._tbl)
    note = anchor.insert_paragraph_before(note_text)
    note.style = anchor.style
    return table, caption, note


def effect(row: dict[str, str], key: str = "HR") -> str:
    return f"{float(row[key]):.2f} ({float(row['lower_95']):.2f}–{float(row['upper_95']):.2f})"


def apply_semantic_styles(document: Document, supplement: bool = False) -> None:
    if "Caption" not in {style.name for style in document.styles}:
        caption_style = document.styles.add_style("Caption", WD_STYLE_TYPE.PARAGRAPH)
        caption_style.base_style = document.styles["Normal"]
        caption_style.font.name = "Times New Roman"
        caption_style.font.size = Pt(9)
        caption_style.font.italic = True
    if document.paragraphs:
        document.paragraphs[0].style = "Title"
    main_level_1 = {
        "Abstract",
        "Introduction",
        "Materials and methods",
        "Results",
        "Discussion",
        "Conclusion",
        "Ethics statement",
        "Author contributions",
        "Funding",
        "Conflict of interest",
        "Data availability statement",
        "References",
    }
    main_level_2 = {
        "Study design and population",
        "Exposure and outcome",
        "Covariates",
        "Causal framework and model definitions",
        "Target-trial-inspired specification and estimands",
        "Statistical analysis",
        "Study population",
        "Main Cox models",
        "Sequential health-status attenuation",
        "Reverse-causation, missing-data, and competing-risk analyses",
        "Regression-standardized absolute risks",
        "Exploratory inflammatory pathway analyses",
        "Secondary penalized and Bayesian analyses",
    }
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if supplement and text.startswith("S") and ". " in text and text[1 : text.index(".")].isdigit():
            paragraph.style = "Heading 1"
        elif not supplement and text in main_level_1:
            paragraph.style = "Heading 1"
        elif not supplement and text in main_level_2:
            paragraph.style = "Heading 2"
        elif text.startswith("Figure ") or text.startswith("Table ") or text.startswith("Supplementary Table "):
            paragraph.style = "Caption"


def build_robustness_table(document: Document) -> None:
    category = read_tsv(BASE / "causal_models_sleep_categories.tsv")
    lag = read_tsv(EXT / "reverse_causation_lag_2y_5y.tsv")
    mi = read_tsv(BASE / "multiple_imputation_rubin_pooled.tsv")
    fg = read_tsv(BASE / "competing_risk_fine_gray.tsv")
    index = {(row["model"], row["term"]): row for row in category + lag + mi + fg}
    specs = [
        ("Free of prevalent CVD", "CVD_free_Model2_confounder", "CVD_free_Model3_health_status", "HR"),
        ("Exclude first 2 years", "lag2y_Model2_confounder", "lag2y_Model3_health_status", "HR"),
        ("Exclude first 5 years", "lag5y_Model2_confounder", "lag5y_Model3_health_status", "HR"),
        ("Multiple imputation", "MI_Model2_confounder", "MI_Model3_health_status", "HR"),
        ("Fine–Gray competing risk", "FineGray_Model2_confounder", "FineGray_Model3_health_status", "SHR"),
    ]
    data = []
    for label, m2, m3, measure in specs:
        r2 = index[(m2, "sleep[>=9 h]")]
        r3 = index[(m3, "sleep[>=9 h]")]
        data.append([label, effect(r2), effect(r3), r2.get("n", ""), r2.get("events", ""), measure])
    table = document.tables[2]
    fill_table(table, ["Analysis", "Model 2", "Model 3", "N", "Heart deaths", "Measure"], data, 7.6)
    replace_prefix(document, "Table 3.", "Table 3. Robustness analyses for long sleep (≥9 versus 7–<8 h/night)")
    replace_prefix(document, "Values are hazard ratios", "Values are effect estimates (95% confidence intervals). Model 2 adjusts the designated baseline confounders; Model 3 additionally adjusts cardiometabolic and pre-existing disease status. HR denotes cause-specific hazard ratio and SHR subdistribution hazard ratio. The 5-year lag sample is smaller because later NHANES cycles cannot contribute participants without at least 5 years of follow-up.")


def absolute_risk_rows():
    risks = read_tsv(EXT / "standardized_absolute_risk_horizons.tsv")
    contrasts = read_tsv(EXT / "standardized_absolute_risk_contrasts.tsv")
    risk_index = {(row["sleep_group"], float(row["horizon_years"])): row for row in risks}
    contrast_index = {(row["sleep_group"], float(row["horizon_years"])): row for row in contrasts}
    output = []
    for level in ("<6 h", "6–<7 h", "7–<8 h", "8–<9 h", "≥9 h"):
        source_level = level.replace("–", "-").replace("≥", ">=")
        r5 = risk_index[(source_level, 5.0)]
        r10 = risk_index[(source_level, 10.0)]
        f5 = f"{100*float(r5['standardized_risk']):.2f}% ({100*float(r5['lower_95_bootstrap']):.2f}–{100*float(r5['upper_95_bootstrap']):.2f})"
        f10 = f"{100*float(r10['standardized_risk']):.2f}% ({100*float(r10['lower_95_bootstrap']):.2f}–{100*float(r10['upper_95_bootstrap']):.2f})"
        if source_level == "7-<8 h":
            rd = "Reference"
        else:
            row = contrast_index[(source_level, 10.0)]
            rd = f"{100*float(row['risk_difference']):.2f} ({100*float(row['rd_lower_95_bootstrap']):.2f}–{100*float(row['rd_upper_95_bootstrap']):.2f})"
        output.append([level, f5, f10, rd])
    return output


def add_main_extensions(document: Document) -> None:
    # Target-trial-inspired framing.
    statistical = find_prefix(document, "Statistical analysis")
    make_heading_before(statistical, "Target-trial-inspired specification and estimands")
    paragraph = statistical.insert_paragraph_before(
        "A target-trial-inspired specification aligned eligibility, exposure classification, and time zero at the NHANES examination [21]. The contrast of interest was habitual sleep ≥9 versus 7–<8 h/night, with follow-up from examination to heart-disease death, competing death, or administrative censoring. The primary estimand was the survey-adjusted cause-specific hazard association under Model 2. A secondary risk-scale estimand was the Model 2 regression-standardized 5- and 10-year cumulative incidence in the presence of competing deaths. These estimands were interpreted associationally because baseline self-reported sleep is not randomized, the intervention is not uniquely defined, and exchangeability cannot be verified."
    )
    paragraph.style = statistical.style

    # Figure 3: sequential attenuation.
    reverse_heading = find_prefix(document, "Reverse-causation, missing-data")
    make_heading_before(reverse_heading, "Sequential health-status attenuation")
    p = reverse_heading.insert_paragraph_before(
        "On the common SIRI-complete sample (n=31,367; 830 heart-disease deaths), the ≥9-h HR was 1.71 (95% CI 1.31–2.25) in the confounder model. Sequential addition of BMI, hypertension/diabetes, prevalent CVD, and log-SIRI reduced it to 1.63, 1.60, 1.54, and 1.50, corresponding to 9.3%, 13.3%, 19.9%, and 25.4% reductions in the log-HR (Figure 3). This descriptive attenuation can arise from confounding, pathway blocking, or non-collapsibility and is not a mediated proportion."
    )
    p.style = reverse_heading.style
    add_picture_before(
        reverse_heading,
        EXT / "Figure6_sequential_health_status_attenuation.png",
        6.5,
        "Forest plot of sequential adjustment stages for the hazard ratio comparing at least 9 with 7 to less than 8 hours of sleep, with attenuation percentages relative to the confounder model.",
    )
    add_caption_before(
        reverse_heading,
        "Figure 3. Sequential health-status adjustment for long sleep. Points and horizontal lines show hazard ratios and 95% confidence intervals for ≥9 versus 7–<8 h/night on the same SIRI-complete sample. Attenuation is the percentage reduction in the log-HR relative to the confounder model and is descriptive rather than a causal mediated proportion.",
        find_prefix(document, "Figure 2."),
    )

    # Figure 4 before Table 3.
    table3 = find_prefix(document, "Table 3.")
    add_picture_before(
        table3,
        EXT / "Figure4_robustness_forest.png",
        6.5,
        "Two-panel forest plot showing the long-sleep association across core sensitivity analyses and after omitting each NHANES cycle in turn.",
    )
    add_caption_before(
        table3,
        "Figure 4. Robustness of the long-sleep association. Panel A compares the confounder and health-status estimates across the main, CVD-free, 2-year-lag, 5-year-lag, multiple-imputation, and Fine–Gray analyses. Panel B omits one NHANES cycle at a time. Fine–Gray rows show subdistribution hazard ratios; all other rows show cause-specific hazard ratios.",
        find_prefix(document, "Figure 2."),
    )

    # Absolute risk section and Figure 5 before inflammation.
    inflammation = find_prefix(document, "Exploratory inflammatory pathway analyses")
    make_heading_before(inflammation, "Regression-standardized absolute risks")
    p = inflammation.insert_paragraph_before(
        "In Model 2, the standardized cumulative incidence of heart-disease death at 10 years was 2.88% (95% PSU-bootstrap interval 2.46%–3.31%) for ≥9 h and 1.87% (1.59%–2.20%) for 7–<8 h. The 10-year risk difference was 1.01 percentage points (0.59–1.49), and the corresponding risk ratio was 1.54 (1.29–1.84). At 5 years, the risk difference was 0.50 percentage points (0.31–0.70) (Table 4 and Figure 5)."
    )
    p.style = inflammation.style
    add_table_before(
        document,
        inflammation,
        "Table 4. Regression-standardized cumulative incidence of heart-disease death",
        ["Sleep duration", "5-year risk", "10-year risk", "10-year RD vs 7–<8 h, percentage points"],
        absolute_risk_rows(),
        "Risks and risk differences are percentages with 95% intervals from 100 stratified PSU bootstrap replicates. Regression standardization used Model 2 cause-specific Cox models for heart-disease and other-cause death and retained competing deaths in the cumulative incidence calculation [23].",
        size=7.5,
    )
    add_picture_before(
        inflammation,
        EXT / "Figure5_standardized_cumulative_incidence.png",
        6.5,
        "Regression-standardized cumulative incidence curves for heart-disease death through 10 years for five sleep-duration groups, including uncertainty bands for the reference and long-sleep groups.",
    )
    add_caption_before(
        inflammation,
        "Figure 5. Model 2 regression-standardized cumulative incidence of heart-disease death in the presence of competing deaths. Shading shows 95% PSU-bootstrap intervals for the 7–<8-h reference and ≥9-h groups. Curves are associational model-based contrasts rather than identified intervention effects.",
        find_prefix(document, "Figure 2."),
    )


def build_main() -> None:
    document = Document(SOURCE_MAIN)
    clean_mojibake(document)
    replacements = {
        "Background: Sleep duration": "Background: Long sleep is associated with cardiovascular mortality, but the signal may reflect underlying illness, frailty, or other unmeasured factors. We tested the association across prespecified causal-model roles and quantified its robustness on relative and absolute risk scales.",
        "Methods: We analyzed adults": "Methods: We analyzed adults aged ≥20 years in NHANES 2005–2018 linked to mortality follow-up. Survey-weighted cause-specific Cox models separated baseline confounders from health-status variables. Extensions included 2- and 5-year lags, leave-one-cycle-out analyses, multiple imputation, Fine–Gray models, sequential health-status adjustment, E-values and deterministic bias scenarios, and regression-standardized cumulative incidence with 100 stratified PSU bootstrap replicates. Skeptical Bayesian shrinkage reported posterior probabilities above clinically interpretable HR thresholds.",
        "Results: The cohort included": "Results: The cohort included 37,516 adults and 1,003 heart-disease deaths; 32,880 participants and 875 deaths formed the common complete-case sample. For ≥9 versus 7–<8 h/night, the HR was 1.70 (95% CI 1.30–2.22) in the main confounder model and 1.53 (1.18–1.99) after health-status adjustment. The 5-year-lag HRs were 2.14 (1.62–2.82) and 1.90 (1.44–2.51), although the lag reduced the sample to 22,610 participants and 446 deaths. Leave-one-cycle-out HRs remained above 1. The Model 2 10-year standardized risk was 2.88% for ≥9 h and 1.87% for 7–<8 h, a difference of 1.01 percentage points (95% bootstrap interval 0.59–1.49). The main HR had an E-value of 2.79 and confidence-limit E-value of 1.92.",
        "Conclusions: Sleep duration showed": "Conclusions: Long sleep was a reproducible risk marker across model definitions, longer lags, survey-cycle exclusions, missing-data, and competing-risk analyses. Health-status attenuation and plausible extreme bias scenarios argue against interpreting the association as a direct effect of sleeping longer. Absolute risks and skeptical posterior probabilities support prognostic relevance while preserving causal caution.",
        "Keywords: sleep duration": "Keywords: sleep duration; heart disease mortality; NHANES; complex survey; causal framework; quantitative bias analysis; regression standardization; competing risks",
        "Long sleep is especially difficult": "Long sleep is especially difficult to interpret because it may reflect poor sleep quality, depression, low physical activity, socioeconomic disadvantage, sleep-disordered breathing, frailty, or occult disease rather than a direct harmful exposure [8, 9, 11–14]. Short sleep has stronger experimental links to sympathetic and inflammatory activation [10], creating a mismatch between the best-established mechanism and the strongest epidemiologic signal. We therefore prioritized an explicit estimand, longer reverse-causation lags, survey-cycle influence checks, quantitative bias analysis, and absolute risks rather than adding prediction models.",
        "This observational cohort study combined": "This observational cohort study combined seven NHANES cycles from 2005–2006 through 2017–2018 and linked participants to public-use mortality follow-up. NHANES uses a multistage probability design to represent the noninstitutionalized US population [15]. Adults with eligible mortality follow-up and sleep duration from 3 to 11 h/night formed the analytic cohort (n=37,516). SIRI and SII were not required for cohort entry because they were exploratory rather than primary adjustment variables. Cross-cycle harmonization and quality-control procedures are detailed in the Supplementary Methods.",
        "Potential variable roles were assigned": "Potential variable roles were assigned before refitting models from the causal structures in Figure 1. Age, sex, race/ethnicity, education, continuous poverty-income ratio (PIR), marital status, and smoking were treated as baseline confounders. BMI, hypertension, diabetes, and prevalent CVD were treated as health-status variables that may reflect pre-existing disease, partial pathways, or both. Prevalent CVD was positive if any of heart failure, coronary heart disease, angina, or myocardial infarction was reported; it was negative only when all four were explicitly absent. SIRI was the primary exploratory inflammatory marker and SII the sensitivity marker. Sleep quality, sleep apnea, depression, physical activity, and frailty remained incompletely measured or unavailable across cycles.",
        "Model 1 adjusted": "Model 1 adjusted for age, sex, and race/ethnicity. Model 2, the main confounder model, additionally adjusted for education, continuous PIR, marital status, and smoking. Model 3, the health-status model, additionally adjusted for continuous BMI, hypertension, diabetes, and prevalent CVD. Age, PIR, and BMI were modeled with four-knot restricted cubic splines at survey-weighted 5th, 35th, 65th, and 95th percentiles. Sequential models added BMI, hypertension/diabetes, CVD, and log-SIRI one domain at a time on a common sample. Elastic net was not used to select confounders or optimize statistical significance.",
        "Baseline characteristics were summarized": "Baseline characteristics were summarized with 14-year MEC weights (WTMEC2YR/7) without null-hypothesis tests. Cause-specific Cox models incorporated MEC weights, SDMVPSU, and SDMVSTRA, with stratified-PSU robust covariance. Models 1–3 were compared in the same complete-case sample. The four nonreference sleep-category coefficients were tested jointly. Sleep splines used knots at 5, 6, 7, 8, and 9 h/night with 7 h/night as the reference. Proportional-hazards estimates were supplemented by standardized cumulative incidence because hazard ratios are conditional and do not directly communicate absolute risk.",
        "Reverse causation was examined": "Reverse causation was examined by excluding prevalent CVD and by excluding deaths and person-time in the first 2 and 5 years. Influence by calendar period was evaluated by leaving out each NHANES cycle in turn. Missing covariates were addressed with 20 stochastic chained-equation imputations and Rubin pooling [16]. Fine–Gray models treated non-heart-disease deaths as competing events [17]. Model 2 standardized risks were obtained by fitting separate cause-specific Cox models for heart-disease and other-cause death, integrating both cause-specific hazards, and averaging predicted risks over the survey-weighted covariate distribution [23]; uncertainty used 100 PSU-with-replacement bootstrap samples within survey strata. E-values were calculated for the point estimate and confidence limit closest to the null, treating the HR as an approximate risk-ratio sensitivity scale given the low absolute risk [22]. Deterministic binary-confounder scenarios varied confounder prevalence in the reference and long-sleep groups and its outcome risk ratio. For pathway triangulation, survey-weighted models related sleep to SIRI/SII; no causal proportion mediated was estimated. Elastic-net and Bayesian shrinkage analyses remained secondary and were not tuned to P values.",
        "On the common sample": "On the common sample, sleep categories were associated with heart-disease mortality in Model 1 (global P<0.001), the main confounder Model 2 (P=0.005), and the health-status Model 3 (P=0.033). Relative to 7–<8 h/night, the ≥9-h HR was 1.70 (95% CI 1.30–2.22) in Model 2 and attenuated to 1.53 (1.18–1.99) in Model 3. In Model 2, estimates for <6 h and 6–<7 h were also elevated, but both attenuated after health-status adjustment; the 8–<9-h interval included 1 in both models.",
        "Additional adjustment": "In the common SIRI-complete sample, sequential addition of BMI, hypertension/diabetes, prevalent CVD, and SIRI reduced the long-sleep log-HR by 9.3%, 13.3%, 19.9%, and 25.4% relative to the confounder model. These changes describe model-dependent attenuation and do not identify a natural or interventional indirect effect.",
        "Among participants free of prevalent CVD": "Among participants free of prevalent CVD, the ≥9-h HR was 1.81 (95% CI 1.28–2.57) in Model 2 and 1.67 (1.19–2.36) in Model 3. After excluding the first 2 years, corresponding HRs were 1.75 (1.33–2.29) and 1.57 (1.20–2.05); after excluding the first 5 years, they were 2.14 (1.62–2.82) and 1.90 (1.44–2.51), based on 22,610 participants and 446 deaths. Because the 5-year lag necessarily removed participants from later cycles with insufficient follow-up, its larger HR should not be read as proof that reverse causation strengthened the association. Across seven leave-one-cycle-out analyses, Model 2 HRs ranged from 1.52 to 1.82 and Model 3 HRs from 1.37 to 1.63, with all confidence intervals excluding 1. Multiple-imputation HRs were 1.73 and 1.58; Fine–Gray SHRs were 1.53 and 1.39 (Figure 4; Table 3). The main Model 2 E-value was 2.79 for the point estimate and 1.92 for its lower confidence limit; Model 3 values were 2.43 and 1.64. Of 32 deterministic bias scenarios, 5 moved the point estimate to or below 1, all requiring pronounced exposure-group imbalance and outcome associations (Supplement).",
        "Sleep duration was associated with SIRI": "Sleep duration was associated with SIRI in categorical (global P=0.008) and spline models (P-overall=0.009; P-nonlinear=0.004), but the magnitude was small: adjusted geometric means were 3.4% higher for <6 h and 3.9% higher for ≥9 h than for 7–<8 h. SII showed no overall spline association and no elevation for ≥9 h. Adding log-SIRI to the sequential common-sample model reduced the long-sleep HR from 1.71 to 1.50 (25.4% reduction in the log-HR across all added health-status domains). This is pathway triangulation, not a mediated proportion, because the marker and exposure were contemporaneous and hazard ratios are non-collapsible.",
        "Figure 3. Survey-weighted": "Figure 6. Survey-weighted restricted cubic splines relating sleep duration to SIRI (Panel A) and SII (Panel B). Values are ratios of adjusted geometric means relative to 7 h/night; shaded regions are 95% confidence intervals. Models adjust the designated baseline confounders. Concurrent measurement of sleep and inflammation precludes a strong causal-mediation interpretation.",
        "The elastic-net one-standard-error": "The elastic-net one-standard-error rule selected no additional laboratory or health-status variable in the primary or repeated PSU-clustered folds. The prediction-optimal lambda-min rule selected eight variables and attenuated the ≥9-h HR on its smaller complete-case sample. Elastic net therefore assessed prediction-oriented selection stability but did not make the etiologic estimate more valid or deliberately more significant. Under the skeptical Normal(0, 0.20²) prior, the health-status posterior HR was 1.34 (95% credible interval 1.08–1.67), with P(HR>1.10)=0.964, P(HR>1.20)=0.844, and P(HR>1.50)=0.159. The declining probability at larger thresholds is more informative than a binary significance claim and does not correct residual confounding.",
        "This reanalysis supports": "This reanalysis supports a nonlinear association between sleep duration and heart-disease mortality, with the most reproducible elevation among adults reporting ≥9 h/night. The association persisted after CVD exclusion, 2- and 5-year lags, imputation, competing-risk modeling, and removal of each survey cycle in turn. Its 10-year standardized risk difference was approximately 1 percentage point. These results establish a robust prognostic marker, not an intervention effect.",
        "The observed shape": "The observed shape is broadly compatible with prior reports of U- or J-shaped sleep–mortality associations [5–7]. However, the estimates for short and moderately long sleep were more sensitive to model definition than the ≥9-h estimate. The 5-year lag estimate increased, but that analysis changed cycle composition and excluded participants who could not accrue 5 years of follow-up. Persistence across lags therefore weakens one simple form of reverse-causation explanation but does not exclude occult disease, frailty, or time-varying health.",
        "Long sleep may be a risk marker": "Long sleep may be a risk marker rather than a modifiable causal exposure. It can accompany fragmented sleep, depression, low activity, unemployment, chronic illness, frailty, or sleep-disordered breathing [8, 9, 11–14]. E-values show that confounding associations of approximately 2.8-fold each would be required to explain the Model 2 point estimate under the E-value framework, but only 1.9-fold each to move its confidence limit to the null. Deterministic scenarios further show that sufficiently strong and imbalanced unmeasured morbidity could remove the association. Neither analysis addresses exposure misclassification or selection bias.",
        "The inflammatory analyses provided": "The inflammatory analyses provided only limited mechanistic support. SIRI showed a shallow nonlinear relation with sleep, whereas SII did not, and sequential health-status adjustment reduced the long-sleep coefficient. Because sleep and blood counts were obtained at the same visit, chronic illness could influence both. The findings are best described as pathway triangulation and health-status attenuation rather than mediation.",
        "Strengths include": "Strengths include a large nationally representative cohort, linked cause-specific mortality, explicit confounder and health-status roles, flexible continuous adjustment, common-sample comparisons, baseline CVD exclusion, 2- and 5-year lags, leave-one-cycle-out analyses, multiple imputation, competing-risk regression, standardized absolute risks, E-values, deterministic bias scenarios, and skeptical posterior threshold probabilities. This hierarchy places design and estimand clarity before prediction-oriented model expansion.",
        "Several limitations": "Several limitations remain substantial. Sleep duration was self-reported once, and sleep quality, obstructive sleep apnea, depression, physical activity, occupational factors, and frailty were incompletely measured or unavailable in the harmonized data. E-values and deterministic scenarios quantify but do not remove unmeasured confounding. The 5-year lag alters calendar-cycle composition, the 100-replicate bootstrap provides only moderately precise tail quantiles, and all standardized risks depend on the two cause-specific proportional-hazards models. Exposure and inflammatory markers were concurrent, precluding temporal identification of an indirect effect. Results apply to the US noninstitutionalized population and should not be interpreted as evidence that changing sleep duration will change mortality risk.",
        "In NHANES": "In NHANES 2005–2018, long sleep was consistently associated with heart-disease mortality across prespecified model roles, longer lags, calendar-cycle exclusions, missing-data, and competing-risk analyses. The absolute 10-year difference was about 1 percentage point. Sequential attenuation, quantitative bias analysis, and weak inflammatory specificity favor interpreting ≥9 h/night as a prognostic marker that may partly reflect underlying poor health. Skeptical Bayesian probabilities quantify directional evidence but do not convert the association into a causal effect.",
        "Original NHANES": "Original NHANES component and linked mortality data are available from the National Center for Health Statistics. Analysis code and tabular outputs for the causal models, longer lags, leave-one-cycle-out analyses, multiple imputation, competing risks, standardized risks, quantitative bias analyses, inflammatory pathway, elastic net, and Bayesian sensitivity analyses accompany this revision.",
    }
    for prefix, text in replacements.items():
        replace_prefix(document, prefix, text)
    build_robustness_table(document)
    add_main_extensions(document)

    ref20 = find_prefix(document, "20.\t")
    refs = [
        "21.\tHernán MA, Sauer BC, Hernández-Díaz S, Platt R, Shrier I: Specifying a target trial prevents immortal time bias and other self-inflicted injuries in observational analyses. J Clin Epidemiol 2016, 79:70–75.",
        "22.\tVanderWeele TJ, Ding P: Sensitivity analysis in observational research: introducing the E-value. Ann Intern Med 2017, 167:268–274.",
        "23.\tSyriopoulou E, Mozumder SI, Rutherford MJ, Lambert PC: Estimating causal effects in the presence of competing events using regression standardisation with the Stata command standsurv. BMC Med Res Methodol 2022, 22:226.",
    ]
    current = ref20
    for text in refs:
        current = insert_paragraph_after(current, text)
        current.style = ref20.style
    clean_mojibake(document)
    document.core_properties.title = "Sleep duration and heart disease mortality among US adults"
    document.core_properties.subject = "Causal framework, quantitative bias analysis, and standardized absolute risks"
    document.core_properties.keywords = "sleep duration; NHANES; heart disease mortality; quantitative bias analysis"
    apply_semantic_styles(document)
    for table, widths in zip(
        document.tables,
        (
            (2000, 1800, 3000, 2500, 900),
            (1600, 2600, 2600, 2600),
            (2100, 1850, 1850, 1200, 1450, 950),
            (1700, 2300, 2300, 3100),
        ),
    ):
        set_table_widths(table, widths)
    document.save(OUTPUT_MAIN)


def add_supp_table(document: Document, anchor: Paragraph, caption: str, headers, rows, note: str = "", size=7.4):
    caption_p = anchor.insert_paragraph_before(caption)
    caption_p.style = anchor.style
    if caption_p.runs:
        caption_p.runs[0].bold = True
    table = document.add_table(rows=1, cols=len(headers))
    table.style = document.tables[0].style
    fill_table(table, headers, rows, size)
    caption_p._p.addnext(table._tbl)
    if note:
        note_p = anchor.insert_paragraph_before(note)
        note_p.style = anchor.style
    return table


def format_supp_effect(row):
    return f"{float(row['HR']):.2f} ({float(row['lower_95']):.2f}–{float(row['upper_95']):.2f})"


def build_supplement() -> None:
    document = Document(SOURCE_SUPPLEMENT)
    clean_mojibake(document)
    replace_prefix(document, "Sleep duration and heart disease mortality", "Sleep duration and heart disease mortality among US adults: advanced causal sensitivity analyses")
    replace_prefix(document, "S2. Variable roles", "S2. Variable roles, target-trial-inspired specification, and estimands")
    s2 = find_prefix(document, "S2.")
    s2_text = insert_paragraph_after(s2, "Eligibility, baseline exposure assignment, and start of follow-up were aligned at the NHANES examination. The primary contrast was ≥9 versus 7–<8 h/night. The primary estimand was the survey-adjusted cause-specific hazard association under the baseline-confounder model; the secondary risk-scale estimand was the regression-standardized 5- and 10-year cumulative incidence in the presence of competing deaths. Neither was interpreted as an identified intervention effect.")
    s2_text.style = s2.style

    replace_prefix(document, "The design-based log-HR", "The design-based log-HR and robust standard error were treated as an approximate normal likelihood and combined with zero-centered normal priors. The skeptical prior was Normal(0, 0.20²) on log(HR), with wider priors examined as sensitivity analyses. Posterior probabilities were reported for HR thresholds of 1, 1.10, 1.20, and 1.50. This is not a full Bayesian survey-survival model and does not model unmeasured confounding. A Bayesian elastic net was not added because penalized selection does not repair the causal identification problem and would make the analysis harder to audit without changing the estimand.")
    bayes = read_tsv(EXT / "bayesian_threshold_probabilities.tsv")
    bayes_rows = []
    for row in bayes:
        bayes_rows.append([
            row["source_model"],
            row["prior"],
            f"{float(row['posterior_HR']):.2f} ({float(row['credible_lower_95']):.2f}–{float(row['credible_upper_95']):.2f})",
            f"{float(row['probability_HR_gt_1']):.3f}",
            f"{float(row['probability_HR_gt_1_10']):.3f}",
            f"{float(row['probability_HR_gt_1_20']):.3f}",
            f"{float(row['probability_HR_gt_1_50']):.3f}",
        ])
    fill_table(document.tables[3], ["Source model", "Prior", "Posterior HR (95% CrI)", "P(HR>1)", "P(HR>1.10)", "P(HR>1.20)", "P(HR>1.50)"], bayes_rows, 6.8)

    reproducibility = find_prefix(document, "S6. Reproducibility files")
    replace_text(reproducibility, "S10. Reproducibility files")
    replace_prefix(document, "Primary code:", "Primary causal code: Data processing/causal_reanalysis_20260808.py. Advanced extension code: Data processing/advanced_causal_extension_20260808.py. Figure code: Data processing/create_causal_figures_20260808.py and Data processing/create_advanced_causal_figures_20260808.py. Validated tabular outputs are in outputs/causal_reanalysis_20260808/ and outputs/causal_extension_20260808/. Elastic-net validation outputs are in outputs/elastic_net_bayesian_sensitivity_20260807/.")

    # S6: lag and cycle robustness.
    make_heading_before(reproducibility, "S6. Longer-lag and calendar-cycle robustness")
    p = reproducibility.insert_paragraph_before("For a k-year lag, participants with follow-up ≤k years were excluded and k years were subtracted from the remaining follow-up; this estimates the post-lag association among those surviving and remaining observable beyond the landmark. The 5-year lag changes the mix of contributing NHANES cycles because later cycles do not have 5 years of potential follow-up. Leave-one-cycle-out analyses refitted both Model 2 and Model 3 seven times.")
    p.style = reproducibility.style
    lag = read_tsv(EXT / "reverse_causation_lag_2y_5y.tsv")
    lag_rows = [[r["model"], format_supp_effect(r), r["n"], r["events"]] for r in lag if r["term"] == "sleep[>=9 h]"]
    add_supp_table(document, reproducibility, "Supplementary Table S5. Longer-lag analyses for ≥9 versus 7–<8 h/night", ["Model", "HR (95% CI)", "N", "Heart deaths"], lag_rows, "Later cycles contribute less or not at all to the 5-year landmark analysis; estimates should be interpreted conditional on post-landmark survival and follow-up availability.")
    loco = read_tsv(EXT / "leave_one_cycle_out_long_sleep.tsv")
    loco_rows = [[r["excluded_cycle"], "Model 2" if r["adjustment_model"] == "Model2_confounder" else "Model 3", format_supp_effect(r), r["n"], r["events"]] for r in loco]
    add_supp_table(document, reproducibility, "Supplementary Table S6. Leave-one-NHANES-cycle-out analyses", ["Excluded cycle", "Adjustment", "HR (95% CI)", "N", "Heart deaths"], loco_rows, "All rows compare ≥9 with 7–<8 h/night and refit the survey-weighted cause-specific Cox model.")

    # S7: sequential attenuation.
    make_heading_before(reproducibility, "S7. Sequential health-status attenuation")
    p = reproducibility.insert_paragraph_before("All stages used the same SIRI-complete sample. Percent attenuation was calculated as 100×(βModel2−βstage)/βModel2. The quantity is descriptive; non-collapsibility and changes in conditioning preclude interpreting it as a mediated proportion.")
    p.style = reproducibility.style
    attenuation = read_tsv(EXT / "health_status_sequential_attenuation.tsv")
    attenuation_rows = []
    for r in attenuation:
        a = "–" if r["attenuation_vs_confounder_percent"] == "nan" else f"{float(r['attenuation_vs_confounder_percent']):.1f}%"
        attenuation_rows.append([r["description"], f"{float(r['HR']):.2f} ({float(r['lower_95']):.2f}–{float(r['upper_95']):.2f})", a, r["n"], r["events"]])
    add_supp_table(document, reproducibility, "Supplementary Table S7. Sequential health-status adjustment", ["Stage", "HR (95% CI)", "Log-HR attenuation", "N", "Deaths"], attenuation_rows)

    # S8: absolute risk.
    make_heading_before(reproducibility, "S8. Regression-standardized absolute risks")
    p = reproducibility.insert_paragraph_before("Two Model 2 survey-weighted cause-specific Cox models were fitted, one for heart-disease death and one for other-cause death. Breslow baseline hazards were combined to obtain individual cumulative incidence functions under each sleep category, which were averaged over the survey-weighted covariate distribution. Uncertainty used PSU-with-replacement bootstrap resampling within survey strata; both models converged in all 100 replicates.")
    p.style = reproducibility.style
    add_supp_table(document, reproducibility, "Supplementary Table S8. Standardized 5- and 10-year risks", ["Sleep duration", "5-year risk", "10-year risk", "10-year RD, percentage points"], absolute_risk_rows(), "Intervals are percentile intervals from 100 stratified PSU bootstrap replicates. These are associational model-based contrasts.")
    contrasts = read_tsv(EXT / "standardized_absolute_risk_contrasts.tsv")
    contrast_rows = []
    for r in contrasts:
        contrast_rows.append([
            r["sleep_group"].replace("-", "–").replace(">=", "≥"),
            f"{float(r['horizon_years']):.0f}",
            f"{100*float(r['risk_difference']):.2f} ({100*float(r['rd_lower_95_bootstrap']):.2f}–{100*float(r['rd_upper_95_bootstrap']):.2f})",
            f"{float(r['risk_ratio']):.2f} ({float(r['rr_lower_95_bootstrap']):.2f}–{float(r['rr_upper_95_bootstrap']):.2f})",
        ])
    add_supp_table(document, reproducibility, "Supplementary Table S9. Standardized risk contrasts versus 7–<8 h/night", ["Sleep duration", "Years", "Risk difference, percentage points", "Risk ratio"], contrast_rows)

    # S9: E-values and deterministic bias scenarios.
    make_heading_before(reproducibility, "S9. Quantitative bias analysis for unmeasured confounding")
    p = reproducibility.insert_paragraph_before("For estimates above 1, E-values used E=HR+√[HR(HR−1)] for the point estimate and confidence limit closest to 1. Because the heart-disease mortality risk was low, the HR was treated as an approximate risk-ratio sensitivity scale. A deterministic binary-confounder analysis additionally varied prevalence in the 7–<8-h group (0.10, 0.20, or 0.30), prevalence in the ≥9-h group (0.30, 0.50, or 0.70), and the confounder–outcome risk ratio (1.5, 2, 3, or 4). The bias factor was [p1RR+(1−p1)]/[p0RR+(1−p0)] and the observed HR was divided by this factor.")
    p.style = reproducibility.style
    evalues = read_tsv(EXT / "e_values_long_sleep.tsv")
    e_rows = [[r["model"], f"{float(r['HR']):.2f}", f"{float(r['lower_95']):.2f}", f"{float(r['E_value_point']):.2f}", f"{float(r['E_value_CI_bound']):.2f}"] for r in evalues]
    add_supp_table(document, reproducibility, "Supplementary Table S10. E-values for long sleep", ["Model", "HR", "Lower 95% limit", "Point E-value", "CI-limit E-value"], e_rows, "An E-value is not evidence that no such confounding exists and does not address selection bias or exposure measurement error.")
    scenarios = read_tsv(EXT / "deterministic_bias_factor_scenarios.tsv")
    scenario_rows = [[f"{float(r['p_unmeasured_confounder_reference']):.2f}", f"{float(r['p_unmeasured_confounder_long_sleep']):.2f}", f"{float(r['RR_unmeasured_confounder_outcome']):.1f}", f"{float(r['bias_factor']):.2f}", f"{float(r['bias_adjusted_HR_approx']):.2f}", "Yes" if r["moves_point_estimate_to_null"] == "True" else "No"] for r in scenarios]
    add_supp_table(document, reproducibility, "Supplementary Table S11. Deterministic binary-confounder scenarios for the main Model 2 estimate", ["p(U), reference", "p(U), ≥9 h", "RR(U–outcome)", "Bias factor", "Adjusted HR", "≤1"], scenario_rows, "The scenario table is deliberately broad. Five of 32 scenarios moved the point estimate to or below 1; each combined a large exposure-group prevalence difference with an outcome RR of 3 or 4.", size=6.7)

    clean_mojibake(document)
    document.core_properties.title = "Supplementary methods and results: advanced causal sensitivity analyses"
    document.core_properties.subject = "NHANES sleep duration and heart-disease mortality"
    apply_semantic_styles(document, supplement=True)
    supplement_widths = (
        (2200, 4300, 3436),
        (2500, 1900, 3000, 2536),
        (3000, 1200, 1200, 4536),
        (2000, 1500, 2100, 1000, 1112, 1112, 1112),
        (3300, 3000, 1800, 1836),
        (2200, 1700, 2500, 1700, 1836),
        (2800, 2600, 1800, 1400, 1336),
        (1700, 2300, 2300, 3636),
        (1700, 1000, 3900, 3336),
        (3000, 1500, 1700, 1900, 1836),
        (1500, 1500, 1700, 1700, 1700, 1836),
    )
    for table, widths in zip(document.tables, supplement_widths):
        set_table_widths(table, widths)
    document.save(OUTPUT_SUPPLEMENT)


def main():
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    build_main()
    build_supplement()
    print(OUTPUT_MAIN)
    print(OUTPUT_SUPPLEMENT)


if __name__ == "__main__":
    main()
