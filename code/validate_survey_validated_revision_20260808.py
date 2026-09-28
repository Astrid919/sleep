"""Validate survey-correction outputs and revised DOCX artifacts."""

from __future__ import annotations

import csv
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn


ROOT = Path(__file__).resolve().parents[1]
BASE = ROOT / "outputs" / "survey_validated_reanalysis_20260808"
EXT = ROOT / "outputs" / "survey_validated_extension_20260808"
DOC = ROOT / "outputs" / "manuscript_survey_validated_20260808"


def read_tsv(path):
    with path.open(encoding="utf-8-sig", newline="") as handle:
        return list(csv.DictReader(handle, delimiter="\t"))


def document_text(document):
    parts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    return "\n".join(parts)


def validate_tables(document):
    for index, table in enumerate(document.tables):
        widths = {len(row.cells) for row in table.rows}
        assert len(widths) == 1, (index, widths)
        assert len(table.rows) >= 1 and len(table.columns) >= 2


def image_alt_text(document):
    output = []
    for element in document.element.body.iter():
        if element.tag == qn("wp:docPr"):
            output.append(element.get("descr", ""))
    return output


def main():
    audit = read_tsv(EXT / "survey_design_audit.tsv")
    assert all(row["status"] == "PASS" for row in audit)
    audit_index = {row["check"]: row["value"] for row in audit}
    assert audit_index["cross_cycle_raw_stratum_collisions"] == "0"
    assert audit_index["cross_cycle_raw_stratum_PSU_collisions"] == "0"
    assert float(audit_index["max_abs_beta_difference_nested_vs_released"]) == 0
    assert float(audit_index["max_abs_covariance_difference_nested_vs_released"]) == 0

    comparison = read_tsv(EXT / "survey_design_estimate_comparison.tsv")
    weighted = next(row for row in comparison if row["model"] == "survey_weighted_explicit_cycle_nesting" and row["term"] == "sleep[>=9 h]")
    unweighted = next(row for row in comparison if row["model"] == "unweighted_explicit_cycle_nesting" and row["term"] == "sleep[>=9 h]")
    assert abs(float(weighted["HR"]) - float(unweighted["HR"])) > 0.05

    ph = read_tsv(EXT / "proportional_hazards_time_interaction.tsv")
    assert all(row["converged"] == "True" for row in ph)
    assert all(0 <= float(row["p_time_interaction"]) <= 1 for row in ph)
    assert all(0 <= float(row["global_sleep_time_interaction_p"]) <= 1 for row in ph)

    same_lag = read_tsv(EXT / "same_cycle_composition_lag_0y_2y_5y.tsv")
    assert {int(row["lag_months"]) for row in same_lag} == {0, 24, 60}
    assert len({row["source_n_before_lag"] for row in same_lag}) == 1

    alternative = read_tsv(EXT / "alternative_4knot_sleep_rcs_summary.tsv")
    assert all(row["knots"] == "5.0,7.0,8.0,9.0" for row in alternative)

    bootstrap = read_tsv(EXT / "standardized_cif_bootstrap_diagnostics.tsv")
    assert len(bootstrap) == 500
    assert all(row["heart_model_converged"] == "True" for row in bootstrap)
    assert all(row["other_model_converged"] == "True" for row in bootstrap)

    mi_spec = read_tsv(BASE / "multiple_imputation_specification.tsv")
    spec_text = " ".join(row["specification"] for row in mi_spec)
    for token in ("multinomial", "Nelson-Aalen", "weight", "stratum", "PSU", "cycle"):
        assert token.lower() in spec_text.lower(), token

    main = Document(DOC / "Manuscript_survey_validated_20260808.docx")
    supplement = Document(DOC / "Supplementary_methods_results_survey_validated_20260808.docx")
    main_text = document_text(main)
    supplement_text = document_text(supplement)
    required_main = (
        "WTMEC2YR/7",
        "cycle × stratum",
        "500 stratified PSU bootstrap",
        "Nelson–Aalen",
        "same 2005–2014 cycle source",
        "robust mortality risk marker",
    )
    for token in required_main:
        assert token in main_text, token
    required_supplement = (
        "weighted Bayesian multinomial logistic",
        "Supplementary Table S11",
        "Supplementary Table S16",
        "Supplementary Figure S1",
        "Supplementary Figure S2",
    )
    for token in required_supplement:
        assert token in supplement_text, token
    forbidden = ("target-trial-inspired", "prespecified causal-model roles", "100 stratified PSU bootstrap", "32 deterministic", "prognostic marker")
    for token in forbidden:
        assert token.lower() not in (main_text + "\n" + supplement_text).lower(), token
    for mojibake in ("≡", "每", "坼", "＊", "�"):
        assert mojibake not in main_text and mojibake not in supplement_text, mojibake
    validate_tables(main)
    validate_tables(supplement)
    main_alt = image_alt_text(main)
    supplement_alt = image_alt_text(supplement)
    assert len(main_alt) >= 6 and all(main_alt)
    assert len(supplement_alt) >= 2 and all(supplement_alt)

    report = [
        {"check": "survey_design_audit", "status": "PASS", "detail": "All pooled-design checks passed."},
        {"check": "weighted_vs_unweighted", "status": "PASS", "detail": f"Weighted HR={float(weighted['HR']):.3f}; unweighted HR={float(unweighted['HR']):.3f}."},
        {"check": "PH_models", "status": "PASS", "detail": "All time-interaction models converged."},
        {"check": "bootstrap", "status": "PASS", "detail": "500/500 heart and other-cause models converged."},
        {"check": "MI_specification", "status": "PASS", "detail": "Education, survival, and survey predictors documented."},
        {"check": "DOCX_content", "status": "PASS", "detail": "Required text present; stale claims absent."},
        {"check": "DOCX_tables_and_alt_text", "status": "PASS", "detail": "Rectangular tables and nonempty image descriptions."},
    ]
    path = DOC / "validation_report.tsv"
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=report[0].keys(), delimiter="\t")
        writer.writeheader()
        writer.writerows(report)
    print(path)


if __name__ == "__main__":
    main()
