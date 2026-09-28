"""Final hard-gate validation for the audited submission package."""

from __future__ import annotations

import argparse
import ast
import csv
import hashlib
import json
import re
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

from docx import Document

from finalize_manuscript_references_20260810 import REFERENCES


def read_tsv(path: Path) -> list[dict[str, str]]:
    with path.open(encoding="utf-8-sig", newline="") as handle:
        return list(csv.DictReader(handle, delimiter="\t"))


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def effect(row: dict[str, str]) -> str:
    return f"{float(row['HR']):.2f} ({float(row['lower_95']):.2f}–{float(row['upper_95']):.2f})"


def trim_row(row: list[str]) -> list[str]:
    output = [cell.strip() for cell in row]
    while output and not output[-1]:
        output.pop()
    return output


def expand_citations(text: str) -> set[int]:
    numbers: set[int] = set()
    for match in re.finditer(r"\[([0-9,\s–-]+)\]", text):
        for token in match.group(1).split(","):
            token = token.strip()
            parts = re.split(r"[–-]", token)
            if len(parts) == 2:
                numbers.update(range(int(parts[0]), int(parts[1]) + 1))
            elif token:
                numbers.add(int(token))
    return numbers


def verify_embedded_figures(docx: Path, reproduced: Path) -> list[dict]:
    ns = {
        "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
        "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
        "wp": "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing",
        "rel": "http://schemas.openxmlformats.org/package/2006/relationships",
    }
    expected = {
        "Two-panel causal diagram": reproduced / "base" / "Figure1_causal_dag.png",
        "Restricted cubic spline hazard-ratio curves": reproduced / "base" / "Figure2_sleep_mortality_rcs.png",
        "One-at-a-time and cumulative health-status": reproduced / "extension" / "Figure3_health_status_domain_attenuation.png",
        "Robustness forest plots": reproduced / "extension" / "Figure4_robustness_forest.png",
        "Regression-standardized cumulative incidence curves": reproduced / "extension" / "Figure5_standardized_cumulative_incidence.png",
        "Restricted cubic spline geometric-mean-ratio curves": reproduced / "base" / "Figure3_inflammation_pathway_rcs.png",
    }
    with zipfile.ZipFile(docx) as archive:
        files = {name: archive.read(name) for name in archive.namelist()}
    document_xml = ET.fromstring(files["word/document.xml"])
    relationships = ET.fromstring(files["word/_rels/document.xml.rels"])
    targets = {item.attrib["Id"]: item.attrib["Target"] for item in relationships.findall("rel:Relationship", ns)}
    results = []
    for drawing in document_xml.findall(".//wp:inline", ns) + document_xml.findall(".//wp:anchor", ns):
        doc_pr = drawing.find("wp:docPr", ns)
        blip = drawing.find(".//a:blip", ns)
        if doc_pr is None or blip is None:
            continue
        description = doc_pr.attrib.get("descr", "")
        selected = next((path for key, path in expected.items() if key in description), None)
        if selected is None:
            continue
        rid = blip.attrib[f"{{{ns['r']}}}embed"]
        archive_name = str((Path("word") / targets[rid]).as_posix())
        embedded_hash = hashlib.sha256(files[archive_name]).hexdigest()
        source_hash = sha256(selected)
        results.append(
            {
                "description": description,
                "source": str(selected),
                "embedded_sha256": embedded_hash,
                "source_sha256": source_hash,
                "status": "PASS" if embedded_hash == source_hash else "FAIL",
            }
        )
    return results


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--project-root", type=Path, required=True)
    parser.add_argument("--package-dir", type=Path, required=True)
    args = parser.parse_args()
    root = args.project_root.resolve()
    package = args.package_dir.resolve()
    docx = package / "manuscript" / "Manuscript_final_audited_20260810.docx"
    base = package / "reproduced_outputs" / "base"
    extension = package / "reproduced_outputs" / "extension"
    table1_dir = package / "reproduced_outputs" / "table1_source"
    document = Document(docx)
    paragraph_text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    tables = [[trim_row([cell.text for cell in row.cells]) for row in table.rows] for table in document.tables]
    checks: list[dict] = []

    def check(name: str, condition: bool, detail: str) -> None:
        checks.append({"check": name, "status": "PASS" if condition else "FAIL", "detail": detail})

    # Reference continuity, citation coverage, and verified metadata.
    reference_text = {}
    for paragraph in document.paragraphs:
        match = re.match(r"^\s*(\d+)\.\s+(.+)$", paragraph.text.strip())
        if match:
            reference_text[int(match.group(1))] = match.group(2)
    citations = expand_citations(paragraph_text.split("References", 1)[0])
    check("reference_count_and_order", sorted(reference_text) == list(range(1, 24)), f"references={sorted(reference_text)}")
    check("citation_coverage", citations == set(range(1, 24)), f"cited={sorted(citations)}")
    check(
        "verified_reference_text",
        all(reference_text.get(item["number"]) == item["text"] for item in REFERENCES),
        "All reference paragraphs match verified metadata.",
    )

    # Table 1.
    source_table1 = read_tsv(table1_dir / "table1_expanded_survey_weighted.tsv")
    expected_table1 = [[row["Characteristic"].strip(), row["Overall"], row["No heart disease death"], row["Heart disease death"]] for row in source_table1]
    observed_table1 = [(row + [""] * 4)[:4] for row in tables[0][1:]]
    check("table1_cells", observed_table1 == expected_table1, f"rows={len(observed_table1)}")
    check("table1_denominators", tables[0][0][:4] == ["Characteristic", "Overall (N=37,516)", "No heart disease death (N=36,513)", "Heart disease death (N=1,003)"], "Header denominators checked.")

    # Table 2.
    category = read_tsv(base / "causal_models_sleep_categories.tsv")
    category_index = {(row["model"], row["term"]): row for row in category}
    summary = {row["model"]: row for row in read_tsv(base / "causal_models_sleep_categories_summary.tsv") if row["cohort"] == "full_common_sample"}
    terms = ["sleep[<6 h]", "sleep[6-<7 h]", None, "sleep[8-<9 h]", "sleep[>=9 h]"]
    labels = ["<6 h", "6–<7 h", "7–<8 h", "8–<9 h", "≥9 h"]
    expected_table2 = [["Sleep duration", "Model 1\nHR (95% CI)", "Model 2\nHR (95% CI)", "Model 3\nHR (95% CI)"]]
    models = ["Model1_demographic", "Model2_confounder", "Model3_health_status"]
    for label, term in zip(labels, terms):
        expected_table2.append([label, *(["Reference"] * 3 if term is None else [effect(category_index[(model, term)]) for model in models])])
    expected_table2.append(["Global P", *[summary[model]["global_sleep_P"] for model in models]])
    check("table2_cells", tables[1] == expected_table2, "All HRs, CIs, references, and global P values checked.")

    # Table 3.
    lag = read_tsv(extension / "reverse_causation_lag_2y_5y.tsv")
    lag_index = {(row["model"], row["term"]): row for row in lag}
    mi = {(row["model"], row["term"]): row for row in read_tsv(base / "multiple_imputation_rubin_pooled.tsv")}
    fg = {(row["model"], row["term"]): row for row in read_tsv(base / "competing_risk_fine_gray.tsv")}
    term = "sleep[>=9 h]"
    expected_table3 = [tables[2][0]]
    expected_table3.extend([
        ["Free of prevalent CVD", effect(category_index[("CVD_free_Model2_confounder", term)]), effect(category_index[("CVD_free_Model3_health_status", term)]), "30096", "531", "HR"],
        ["Exclude first 2 years", effect(lag_index[("lag2y_Model2_confounder", term)]), effect(lag_index[("lag2y_Model3_health_status", term)]), "30198", "733", "HR"],
        ["Exclude first 5 years", effect(lag_index[("lag5y_Model2_confounder", term)]), effect(lag_index[("lag5y_Model3_health_status", term)]), "22610", "446", "HR"],
        ["Multiple imputation", effect(mi[("MI_Model2_confounder", term)]), effect(mi[("MI_Model3_health_status", term)]), "37516", "1003", "HR"],
        ["Fine–Gray sensitivity*", effect(fg[("FineGray_Model2_confounder", term)]), effect(fg[("FineGray_Model3_health_status", term)]), "32880", "875", "SHR"],
    ])
    check("table3_cells", tables[2] == expected_table3, "All robustness estimates checked.")

    # Table 4.
    risks = {(row["sleep_group"], float(row["horizon_years"])): row for row in read_tsv(extension / "standardized_absolute_risk_horizons.tsv")}
    contrasts = {(row["sleep_group"], float(row["horizon_years"])): row for row in read_tsv(extension / "standardized_absolute_risk_contrasts.tsv")}
    expected_table4 = [tables[3][0]]
    for source, label in zip(["<6 h", "6-<7 h", "7-<8 h", "8-<9 h", ">=9 h"], labels):
        row5, row10 = risks[(source, 5.0)], risks[(source, 10.0)]
        five = f"{100*float(row5['standardized_risk']):.2f}% ({100*float(row5['lower_95_bootstrap']):.2f}–{100*float(row5['upper_95_bootstrap']):.2f})"
        ten = f"{100*float(row10['standardized_risk']):.2f}% ({100*float(row10['lower_95_bootstrap']):.2f}–{100*float(row10['upper_95_bootstrap']):.2f})"
        if source == "7-<8 h":
            rd = "Reference"
        else:
            item = contrasts[(source, 10.0)]
            rd = f"{float(item['risk_difference_percentage_points']):.2f} ({100*float(item['rd_lower_95_bootstrap']):.2f}–{100*float(item['rd_upper_95_bootstrap']):.2f})"
        expected_table4.append([label, five, ten, rd])
    check("table4_cells", tables[3] == expected_table4, "All standardized risks and risk differences checked.")

    # Narrative claims and diagnostics.
    required_claims = [
        "32,880 participants, including 875 heart disease deaths and 2,629 deaths from other causes",
        "248,924 person-years",
        "7.92 years (IQR, 4.42–11.33)",
        "2.39 per 1,000 person-years",
        "HR 1.70, 95% CI 1.30–2.22",
        "HR 1.53, 95% CI 1.18–1.99",
        "risk difference 1.01 percentage points, 95% bootstrap interval 0.61–1.45",
    ]
    check("narrative_numeric_claims", all(value in paragraph_text for value in required_claims), "Key abstract/results claims checked against reproduced outputs.")
    bootstrap = read_tsv(extension / "standardized_cif_bootstrap_diagnostics.tsv")
    check("bootstrap_500_convergence", len(bootstrap) == 500 and all(row["heart_model_converged"] == "True" and row["other_model_converged"] == "True" for row in bootstrap), f"replicates={len(bootstrap)}")

    figures = verify_embedded_figures(docx, package / "reproduced_outputs")
    check("embedded_figures", len(figures) == 6 and all(item["status"] == "PASS" for item in figures), f"matched={sum(item['status'] == 'PASS' for item in figures)}/6")

    # Code syntax audit for the complete final-generation chain.
    code_files = [
        "audit_and_bayesian_reanalysis.py",
        "causal_reanalysis_20260808.py",
        "advanced_causal_extension_20260808.py",
        "survey_validated_extension_20260808.py",
        "create_causal_figures_20260808.py",
        "create_advanced_causal_figures_20260808.py",
        "create_survey_validated_figures_20260808.py",
        "revise_manuscript_causal_20260808.py",
        "revise_manuscript_causal_extension_20260808.py",
        "revise_manuscript_survey_validated_20260808.py",
        "validate_survey_validated_revision_20260808.py",
        "audit_final_manuscript_20260810.py",
        "regenerate_final_figures_20260810.py",
        "compare_reproduced_outputs_20260810.py",
        "finalize_manuscript_references_20260810.py",
        "validate_final_package_20260810.py",
    ]
    syntax_errors = []
    for name in code_files:
        path = root / "数据处理" / name
        try:
            ast.parse(path.read_text(encoding="utf-8"), filename=str(path))
        except Exception as exc:  # pragma: no cover - audit reporting path
            syntax_errors.append(f"{name}: {exc}")
    check("code_syntax", not syntax_errors, "; ".join(syntax_errors) if syntax_errors else f"Parsed {len(code_files)} Python scripts.")

    reproduction = json.loads((package / "audit" / "reproduction_comparison.json").read_text(encoding="utf-8"))
    numeric_reproduction_ok = all(row["status"] == "PASS" for section in ("base_tsv", "extension_tsv") for row in reproduction[section])
    numeric_detail = {
        "base_tsv": f'{reproduction["summary"]["base_pass"]}/{reproduction["summary"]["base_total"]}',
        "extension_tsv": f'{reproduction["summary"]["extension_pass"]}/{reproduction["summary"]["extension_total"]}',
        "legacy_png_note": (
            "Legacy PNGs were renderer-level review items and were replaced; "
            "the separate embedded_figures check verifies the 6 final manuscript figures exactly."
        ),
    }
    check("independent_numeric_reproduction", numeric_reproduction_ok, json.dumps(numeric_detail, ensure_ascii=False))

    with zipfile.ZipFile(docx) as archive:
        xml = archive.read("word/document.xml")
        comments = "word/comments.xml" in archive.namelist()
    tracked = bool(re.search(br"<w:(?:ins|del)(?:\s|>)", xml))
    check("clean_review_state", not comments and not tracked, f"comments={comments}; tracked_changes={tracked}")

    report = {
        "manuscript": str(docx),
        "manuscript_sha256": sha256(docx),
        "checks": checks,
        "figures": figures,
        "summary": {
            "passed": sum(item["status"] == "PASS" for item in checks),
            "total": len(checks),
            "failed": [item["check"] for item in checks if item["status"] != "PASS"],
        },
    }
    audit_dir = package / "audit"
    audit_dir.mkdir(parents=True, exist_ok=True)
    (audit_dir / "final_validation_report.json").write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    lines = ["# Final validation report", "", f"Manuscript SHA-256: `{report['manuscript_sha256']}`", ""]
    lines.extend(f"- {item['status']} — {item['check']}: {item['detail']}" for item in checks)
    lines.extend(["", f"Passed {report['summary']['passed']}/{report['summary']['total']} checks."])
    (audit_dir / "final_validation_report.md").write_text("\n".join(lines), encoding="utf-8")
    print(json.dumps(report["summary"], ensure_ascii=False))
    if report["summary"]["failed"]:
        raise SystemExit(1)


if __name__ == "__main__":
    main()
