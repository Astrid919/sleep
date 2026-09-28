"""Cross-check final manuscript structure and reported estimates against TSVs."""

from __future__ import annotations

import csv
import json
from pathlib import Path

from docx import Document


ROOT = Path(__file__).resolve().parents[1]
RESULTS = ROOT / "outputs" / "causal_reanalysis_20260808"
DOCX = ROOT / "outputs" / "manuscript_causal_revision_20260808" / "Manuscript_causal_revision_20260808.docx"
REPORT = ROOT / "outputs" / "manuscript_causal_revision_20260808" / "validation_report_20260808.json"


def rows(name: str):
    with (RESULTS / name).open(encoding="utf-8-sig", newline="") as handle:
        return list(csv.DictReader(handle, delimiter="\t"))


def find(data, model, term):
    matches = [row for row in data if row.get("model") == model and row.get("term") == term]
    assert len(matches) == 1, (model, term, len(matches))
    return matches[0]


def rounded_effect(row):
    return tuple(round(float(row[name]), 2) for name in ("HR", "lower_95", "upper_95"))


def main():
    category = rows("causal_models_sleep_categories.tsv")
    mi = rows("multiple_imputation_rubin_pooled.tsv")
    fine_gray = rows("competing_risk_fine_gray.tsv")
    lag = rows("reverse_causation_lag_24_months.tsv")
    qc = {row["metric"]: int(row["value"]) for row in rows("causal_reanalysis_quality_control.tsv")}
    expected = {
        "complete_model2_long": rounded_effect(find(category, "Model2_confounder", "sleep[>=9 h]")),
        "complete_model3_long": rounded_effect(find(category, "Model3_health_status", "sleep[>=9 h]")),
        "cvdfree_model2_long": rounded_effect(find(category, "CVD_free_Model2_confounder", "sleep[>=9 h]")),
        "cvdfree_model3_long": rounded_effect(find(category, "CVD_free_Model3_health_status", "sleep[>=9 h]")),
        "mi_model2_long": rounded_effect(find(mi, "MI_Model2_confounder", "sleep[>=9 h]")),
        "mi_model3_long": rounded_effect(find(mi, "MI_Model3_health_status", "sleep[>=9 h]")),
        "fg_model2_long": rounded_effect(find(fine_gray, "FineGray_Model2_confounder", "sleep[>=9 h]")),
        "fg_model3_long": rounded_effect(find(fine_gray, "FineGray_Model3_health_status", "sleep[>=9 h]")),
        "lag_model2_long": rounded_effect(find(lag, "lag24_Model2_confounder", "sleep[>=9 h]")),
        "lag_model3_long": rounded_effect(find(lag, "lag24_Model3_health_status", "sleep[>=9 h]")),
    }
    required = {
        "complete_model2_long": (1.70, 1.30, 2.22),
        "complete_model3_long": (1.53, 1.18, 1.99),
        "cvdfree_model2_long": (1.81, 1.28, 2.57),
        "cvdfree_model3_long": (1.67, 1.19, 2.36),
        "mi_model2_long": (1.73, 1.35, 2.23),
        "mi_model3_long": (1.58, 1.24, 2.02),
        "fg_model2_long": (1.53, 1.17, 2.02),
        "fg_model3_long": (1.39, 1.06, 1.83),
        "lag_model2_long": (1.75, 1.33, 2.29),
        "lag_model3_long": (1.57, 1.20, 2.05),
    }
    assert expected == required, {"actual": expected, "required": required}
    assert qc["expanded_cohort_n"] == 37516
    assert qc["expanded_cohort_heart_deaths"] == 1003
    assert qc["common_complete_case_n"] == 32880
    assert qc["common_complete_case_heart_deaths"] == 875
    assert qc["strict_CVD_free_complete_case_n"] == 30096
    assert qc["strict_CVD_free_complete_case_heart_deaths"] == 531
    assert qc["multiple_imputations"] == 20

    document = Document(DOCX)
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    assert len(document.inline_shapes) == 3
    assert len(document.tables) == 3
    assert [(len(table.rows), len(table.columns)) for table in document.tables] == [(47, 5), (7, 4), (7, 6)]
    for phrase in (
        "Covariate selection with LASSO",
        "The rebuilt cohort included",
        "missing-category sensitivity",
        "fully adjusted causal model",
        "was mediated by SIRI",
    ):
        assert phrase not in text, phrase
    for phrase in (
        "Causal framework and model definitions",
        "Reverse-causation, missing-data, and competing-risk analyses",
        "Exploratory inflammatory pathway analyses",
        "Secondary penalized and Bayesian analyses",
        "30,096 participants free of CVD",
        "20 stochastic chained-equation imputations",
    ):
        assert phrase in text, phrase
    alt_texts = [shape._inline.docPr.get("descr") for shape in document.inline_shapes]
    assert all(alt_texts)
    result = {
        "status": "PASS",
        "validated_effects": expected,
        "sample_counts": qc,
        "docx": {
            "paragraphs": len(document.paragraphs),
            "tables": len(document.tables),
            "inline_images": len(document.inline_shapes),
            "alt_texts": alt_texts,
        },
    }
    REPORT.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    print(REPORT)


if __name__ == "__main__":
    main()
