"""Validate the advanced analysis outputs and final manuscript artifacts."""

from __future__ import annotations

import csv
import json
import math
from pathlib import Path

from docx import Document


ROOT = Path(__file__).resolve().parents[1]
BASE = ROOT / "outputs" / "causal_reanalysis_20260808"
EXT = ROOT / "outputs" / "causal_extension_20260808"
DOC_DIR = ROOT / "outputs" / "manuscript_causal_extension_20260808"
MAIN = DOC_DIR / "Manuscript_causal_extension_20260808.docx"
SUPPLEMENT = DOC_DIR / "Supplementary_methods_results_causal_extension_20260808.docx"
REPORT = DOC_DIR / "validation_report_causal_extension_20260808.json"


def rows(path: Path):
    with path.open(encoding="utf-8-sig", newline="") as handle:
        return list(csv.DictReader(handle, delimiter="\t"))


def row_for(data, **criteria):
    matches = [row for row in data if all(row.get(key) == value for key, value in criteria.items())]
    assert len(matches) == 1, (criteria, len(matches))
    return matches[0]


def close(actual, expected, tolerance=5e-6):
    assert math.isclose(float(actual), expected, rel_tol=0, abs_tol=tolerance), (actual, expected)


def document_text(document: Document):
    return "\n".join(paragraph.text for paragraph in document.paragraphs) + "\n" + "\n".join(
        cell.text for table in document.tables for row in table.rows for cell in row.cells
    )


def main():
    lag = rows(EXT / "reverse_causation_lag_2y_5y.tsv")
    attenuation = rows(EXT / "health_status_sequential_attenuation.tsv")
    loco = rows(EXT / "leave_one_cycle_out_long_sleep.tsv")
    evalues = rows(EXT / "e_values_long_sleep.tsv")
    scenarios = rows(EXT / "deterministic_bias_factor_scenarios.tsv")
    risks = rows(EXT / "standardized_absolute_risk_horizons.tsv")
    contrasts = rows(EXT / "standardized_absolute_risk_contrasts.tsv")
    diagnostics = rows(EXT / "standardized_cif_bootstrap_diagnostics.tsv")
    bayes = rows(EXT / "bayesian_threshold_probabilities.tsv")

    expected_lag = {
        "lag2y_Model2_confounder": (1.74513347, 1.33146321, 2.28732631, 30198, 733),
        "lag2y_Model3_health_status": (1.56738076, 1.19887570, 2.04915526, 30198, 733),
        "lag5y_Model2_confounder": (2.13657843, 1.61703211, 2.82305302, 22610, 446),
        "lag5y_Model3_health_status": (1.90112260, 1.43972968, 2.51037897, 22610, 446),
    }
    for model, expected in expected_lag.items():
        row = row_for(lag, model=model, term="sleep[>=9 h]")
        for key, value in zip(("HR", "lower_95", "upper_95"), expected[:3]):
            close(row[key], value)
        assert int(row["n"]) == expected[3]
        assert int(row["events"]) == expected[4]

    stage_b = row_for(attenuation, stage="B_confounder")
    stage_f = row_for(attenuation, stage="F_plus_SIRI")
    close(stage_b["HR"], 1.71421796)
    close(stage_f["HR"], 1.49503760)
    close(stage_f["attenuation_vs_confounder_percent"], 25.38340228)
    assert len(loco) == 14
    for model, expected_range in {
        "Model2_confounder": (1.51832020, 1.82447101),
        "Model3_health_status": (1.37384870, 1.62963280),
    }.items():
        values = [float(row["HR"]) for row in loco if row["adjustment_model"] == model]
        close(min(values), expected_range[0])
        close(max(values), expected_range[1])
        assert all(float(row["lower_95"]) > 1 for row in loco if row["adjustment_model"] == model)

    main_e = row_for(evalues, model="Model2_confounder")
    close(main_e["E_value_point"], 2.78612638)
    close(main_e["E_value_CI_bound"], 1.91855769)
    assert len(scenarios) == 32
    assert sum(row["moves_point_estimate_to_null"] == "True" for row in scenarios) == 5

    risk_ref = row_for(risks, sleep_group="7-<8 h", horizon_years="10.0")
    risk_long = row_for(risks, sleep_group=">=9 h", horizon_years="10.0")
    close(risk_ref["risk_percent"], 1.87128465)
    close(risk_long["risk_percent"], 2.88269471)
    contrast = row_for(contrasts, sleep_group=">=9 h", horizon_years="10.0")
    close(contrast["risk_difference_percentage_points"], 1.01141006)
    close(contrast["rd_lower_95_bootstrap"], 0.00594605)
    close(contrast["rd_upper_95_bootstrap"], 0.01485975)
    assert len(diagnostics) == 100
    assert all(row["heart_model_converged"] == "True" and row["other_model_converged"] == "True" for row in diagnostics)
    for row in risks:
        value = float(row["standardized_risk"])
        assert float(row["lower_95_bootstrap"]) <= value <= float(row["upper_95_bootstrap"])

    skeptical = row_for(bayes, source_model="Model3_health_status", prior="Normal(0,0.2^2)")
    close(skeptical["posterior_HR"], 1.34254606)
    close(skeptical["probability_HR_gt_1_10"], 0.96387609)
    close(skeptical["probability_HR_gt_1_20"], 0.84437483)
    close(skeptical["probability_HR_gt_1_50"], 0.15855079)

    main_doc = Document(MAIN)
    supp_doc = Document(SUPPLEMENT)
    main_text = document_text(main_doc)
    supp_text = document_text(supp_doc)
    assert len(main_doc.inline_shapes) == 6
    assert all(shape._inline.docPr.get("descr") for shape in main_doc.inline_shapes)
    assert [(len(table.rows), len(table.columns)) for table in main_doc.tables] == [(47, 5), (7, 4), (6, 6), (6, 4)]
    assert len(supp_doc.tables) == 11
    assert [(len(table.rows), len(table.columns)) for table in supp_doc.tables[3:]] == [
        (13, 7), (5, 4), (15, 5), (7, 5), (6, 4), (9, 4), (11, 5), (33, 6)
    ]
    required_main = (
        "Target-trial-inspired specification and estimands",
        "Sequential health-status attenuation",
        "Regression-standardized absolute risks",
        "E-value was 2.79",
        "P(HR>1.50)=0.159",
        "Figure 6.",
        "23.\tSyriopoulou",
    )
    for phrase in required_main:
        assert phrase in main_text, phrase
    required_supp = (
        "S6. Longer-lag and calendar-cycle robustness",
        "S9. Quantitative bias analysis",
        "A Bayesian elastic net was not added",
        "Supplementary Table S11.",
    )
    for phrase in required_supp:
        assert phrase in supp_text, phrase
    for bad in ("每", "≡", "坼", "＊", "杅擂揭燴", "�"):
        assert bad not in main_text + supp_text, bad
    for phrase in (
        "was mediated by SIRI",
        "fully adjusted causal model",
        "causal effect of sleeping longer",
    ):
        assert phrase not in main_text.lower(), phrase

    report = {
        "status": "PASS",
        "analysis": {
            "lag_models": expected_lag,
            "leave_one_cycle_out_rows": len(loco),
            "deterministic_bias_scenarios": len(scenarios),
            "bias_scenarios_moving_point_to_null": 5,
            "bootstrap_replicates": len(diagnostics),
            "bootstrap_models_converged": 100,
        },
        "main_docx": {
            "paragraphs": len(main_doc.paragraphs),
            "tables": len(main_doc.tables),
            "inline_images": len(main_doc.inline_shapes),
            "all_images_have_alt_text": True,
        },
        "supplement_docx": {
            "paragraphs": len(supp_doc.paragraphs),
            "tables": len(supp_doc.tables),
        },
    }
    REPORT.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(REPORT)


if __name__ == "__main__":
    main()
